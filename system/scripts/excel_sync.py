"""creaite Excel round-trip — tracker-view.xlsx, the Excel editing surface.

Commands:
  export    write tracker-view.xlsx from tracker.db (one sheet per table, plus a Review sheet).
            Refuses to overwrite a workbook with unimported edits.
  import    stage the workbook's differences as a proposal (source_run: excel-import), and copy
            Review sheet decisions into the pending proposals.
  changed   report whether the workbook differs from the last export (exit 1 if it has unimported edits)

The workbook is never the record: everything it changes goes through system/staging/ and tracker.py apply.
Dates are real Excel dates shown DD-MM-YYYY; they are stored YYYY-MM-DD.
"""
import hashlib
import json
import sqlite3
import subprocess
import sys
from datetime import date, datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import paths  # noqa: E402  (puts system/scripts/vendor/ first)
try:
    import log as debuglog  # BL-102: the debug log (not `log`: some scripts have a function of that name); optional
except ImportError:
    class debuglog:  # the helper is missing: every call does nothing and the script runs as before
        info = debug = warn = error = exception = staticmethod(lambda *a, **k: None)
        rule = staticmethod(lambda text, *a: "")
        main = staticmethod(lambda script, fn, *a: fn(*a))

import review  # noqa: E402  (shared review details, BL-25)
import tracker  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

STATE = paths.STAGING / "export-state.json"
LOCK = paths.EXCEL_VIEW.with_name("~$" + paths.EXCEL_VIEW.name)
EXTRA_ROWS = 500          # rows below the data that keep drop-downs and formats for new entries

# Sheet layout: (sheet name, table, key fields, [(column label, field)], locked fields)
SHEETS = [
    ("Projects", "projects", ("code",), [
        ("Code", "code"), ("Name", "name"), ("Sponsor", "sponsor"), ("Phase", "phase"), ("Health", "health"),
        ("Health reason", "health_reason"), ("Next milestone", "next_milestone"),
        ("Milestone date", "milestone_date"), ("Active", "active"), ("Updated", "updated_at")],
     {"code", "updated_at"}),
    ("Milestones", "milestones", ("id",), [
        ("ID", "id"), ("Project", "project"), ("Milestone", "name"), ("Description", "description"), ("Baseline", "baseline_date"),
        ("Forecast", "forecast_date"), ("Status", "status"), ("Order", "sort_order"), ("% complete", "percent_complete"),
        ("Owner", "owner"), ("State", "state"), ("State set by you", "state_set"), ("State reason", "state_reason"),   # BL-186
        ("Achieved on", "actual_date"), ("Client-visible", "client_visible"), ("Client wording", "client_message"),
        ("Internal note", "internal_note"), ("Source", "source"), ("Updated", "updated_at")],
     {"id", "source", "updated_at"}),
    ("Workstreams", "workstreams", ("id",), [                                  # BL-58: editable like the others
        ("ID", "id"), ("Project", "project"), ("Workstream", "name"), ("Lead", "lead"), ("Keywords", "keywords"),
        ("Aliases", "aliases"), ("Teams channel", "channel"), ("Status", "status"), ("Created", "created"),
        ("Source", "source"), ("Updated", "updated_at")],
     {"id", "created", "source", "updated_at"}),
    ("Actions", "actions", ("id",), [
        ("ID", "id"), ("Project", "project"), ("Description", "description"), ("Workstream", "workstream"),
        ("Owner", "owner"), ("Priority", "priority"), ("Due", "due_date"), ("Status", "status"),
        ("Raised", "raised"), ("Closed", "closed"), ("Latest update", "latest_update"),
        ("Client-visible", "client_visible"), ("Client messaging", "client_message"), ("Milestone", "milestone"),
        ("Labels", "labels"), ("Source", "source"), ("Updated", "updated_at")],
     {"id", "source", "updated_at"}),
    ("RAID", "raid", ("id",), [
        ("ID", "id"), ("Project", "project"), ("Type", "type"), ("Workstream", "workstream"),
        ("Description", "description"), ("Owner", "owner"), ("Response action", "response_action"),
        ("Due", "due_date"), ("Impact", "impact"), ("Likelihood", "likelihood"), ("Priority", "priority"),
        ("Status", "status"), ("Raised", "raised"), ("Closed", "closed"), ("Client-visible", "client_visible"),
        ("Client messaging", "client_message"), ("Milestone", "milestone"), ("Labels", "labels"), ("Source", "source"),
        ("Updated", "updated_at")],
     {"id", "priority", "source", "updated_at"}),
    ("Decisions", "decisions", ("id",), [                                      # BL-53: status, needed by, reasoning, options, revisit, replaces
        ("ID", "id"), ("Project", "project"), ("Decision", "decision"), ("Status", "status"), ("Made by", "made_by"),
        ("Needed by", "needed_by"), ("Date", "date"), ("Reasoning", "reasoning"), ("Options", "options"),
        ("Revisit", "revisit"), ("Replaces", "replaces"), ("Client-visible", "client_visible"),
        ("Client messaging", "client_message"), ("Labels", "labels"), ("Source", "source"), ("Updated", "updated_at")],
     {"id", "source", "updated_at"}),
    ("Promises", "promises", ("id",), [                                        # BL-50: who promised what to whom, by when
        ("ID", "id"), ("Project", "project"), ("Promised by", "promised_by"), ("Promised to", "promised_to"),
        ("Promise", "description"), ("Due", "due_date"), ("Status", "status"), ("Raised", "raised"), ("Closed", "closed"),
        ("Source", "source"), ("Updated", "updated_at")],
     {"id", "closed", "source", "updated_at"}),
    ("Links", "links", ("from_item", "to_item", "kind"), [                    # BL-90 part 2: read-only in Excel at first
        ("Blocker", "from_item"), ("Blocks", "to_item"), ("Kind", "kind"), ("Created", "created"), ("Source", "source")],
     {"from_item", "to_item", "kind", "created", "source"}),
    ("Updates", "updates", ("week_ending", "project"), [
        ("Week ending", "week_ending"), ("Project", "project"), ("Health", "health"), ("Summary", "summary"),
        ("Key changes", "key_changes")],
     {"week_ending", "project"}),
]
DATES = {"milestone_date", "due_date", "raised", "closed", "date", "week_ending", "baseline_date", "forecast_date",
         "created", "needed_by", "actual_date"}
READ_ONLY = {"links"}       # sheets Excel only shows; edits to them are ignored on import (the sheet says so)
LABELLED = tracker.LABELLED # tables whose rows carry a Labels column (a comma-separated list; never a drop-down)
YES_NO = {"client_visible", "active", "state_set"}
WIDE = {"description", "decision", "client_message", "latest_update", "response_action", "health_reason",
        "summary", "key_changes", "source", "reasoning", "options", "revisit", "state_reason", "internal_note"}
CHOICES = {
    "health": ["On track", "At risk", "Off track"],
    ("actions", "status"): ["Open", "Delayed", "Closed", "Cancelled"],
    ("raid", "status"): ["Open", "Closed", "Cancelled"],
    ("milestones", "status"): ["Planned", "Done", "Cancelled"],              # BL-186: Missed is retired
    ("milestones", "state"): ["On track", "At risk", "Slipped", "Overdue", "Achieved"],
    ("workstreams", "status"): ["Active", "Closed"],
    ("decisions", "status"): ["Pending", "Made", "Withdrawn"],
    ("promises", "status"): ["Open", "Done", "Dropped"],
    "priority": ["Urgent", "High", "Medium", "Low"],
    "type": ["Risk", "Assumption", "Issue", "Dependency"],
    "impact": ["High", "Medium", "Low"],
    "likelihood": ["High", "Medium", "Low"],
}
REVIEW_COLS = [("File", 34), ("Item", 7), ("Project", 9), ("What", 16), ("ID", 12), ("Proposed", 70),
               ("Flags", 18), ("Match", 9), ("Age (days)", 10), ("Decision", 12), ("Note", 30),
               ("From minutes", 34)]   # BL-41: the draft minutes an item came from (kept last: J and K stay put)

HEADER = PatternFill("solid", fgColor="DCE6E4")
LOCKED = PatternFill("solid", fgColor="EFEFEF")


# ── Helpers ───────────────────────────────────────────────────────────────────

def db():
    """A read-only connection to tracker.db, rows by column name; exits with a message if it is missing."""
    if not paths.DB.exists():
        sys.exit(f"creaite: {paths.rel(paths.DB)} not found. Run: python3 system/scripts/tracker.py init")
    con = sqlite3.connect(f"file:{paths.DB}?mode=ro", uri=True)
    con.row_factory = sqlite3.Row
    return con


def file_hash():
    """The SHA-256 of tracker-view.xlsx, or None when there is no workbook."""
    return hashlib.sha256(paths.EXCEL_VIEW.read_bytes()).hexdigest() if paths.EXCEL_VIEW.exists() else None


def load_state():
    """The last export's record (system/staging/export-state.json): its rows and file hash."""
    return json.loads(STATE.read_text(encoding="utf-8")) if STATE.exists() else {}


def save_state(state):
    """Write the export record to system/staging/export-state.json."""
    STATE.parent.mkdir(parents=True, exist_ok=True)
    STATE.write_text(json.dumps(state, indent=1, ensure_ascii=False), encoding="utf-8")


def has_unimported_edits():
    """True if the workbook exists and differs from both the last export and the last import."""
    h = file_hash()
    if h is None:
        return False
    state = load_state()
    return h not in (state.get("hash"), state.get("imported_hash"))


def key_of(row, keys):
    """A row's key as one string: the values of `keys` joined with '|'."""
    return "|".join(str(row[k]) for k in keys)


def choices(table, field):
    """The drop-down values for a field, for this table first, then for the field anywhere."""
    return CHOICES.get((table, field)) or CHOICES.get(field)


def to_cell(field, value):
    """A stored value as it goes in a cell: dates as real dates, flags as Yes or No."""
    if value is None:
        return None
    if field in DATES:
        try:
            return date.fromisoformat(value)
        except (TypeError, ValueError):
            return value
    if field in YES_NO:
        return "Yes" if int(value) else "No"
    return value


def from_cell(field, value):
    """Workbook value → stored value (dates YYYY-MM-DD, Yes/No → 1/0, blanks → None)."""
    if value is None:
        return None
    if isinstance(value, datetime):
        value = value.date()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str):
        if value.strip() == "":
            return None
        if field in DATES:
            try:
                return datetime.strptime(value.strip(), "%d-%m-%Y").date().isoformat()
            except ValueError:
                return value  # left as typed; tracker.py validate reports it
    if field in YES_NO:
        if str(value).lower() in ("yes", "y", "1", "true"):
            return 1
        if str(value).lower() in ("no", "n", "0", "false"):
            return 0
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def same(a, b, field=None):
    """Equal once blanks and trailing spaces are ignored (Excel may drop them); labels compare as a set, in any order and case."""
    if field == "labels":
        return {n.lower() for n in tracker.label_names(a)} == {n.lower() for n in tracker.label_names(b)}

    def norm(x):
        """The value with trailing spaces dropped; blank is None."""
        return None if x is None or str(x).strip() == "" else "\n".join(
            line.rstrip() for line in str(x).rstrip().splitlines())
    return norm(a) == norm(b)


def blank_default(table, field):
    """What a cleared cell means for fields that can't be empty."""
    if field == "client_message":
        return "" if table == "decisions" else (
            "What it is: \nWhat it means for you: \nWhat is being done: ")
    return None


# ── Export ────────────────────────────────────────────────────────────────────

def shown_flags(item):
    """A proposal's flags as the Review sheet shows them. BL-288: `vip` (who is one is private to Alan) is never shown."""
    return ", ".join(f for f in item.get("flags", []) if f != "vip")


def pending_rows():
    """Every pending item and draft block in system/staging/, for the Review sheet."""
    out = []
    for p in sorted(paths.STAGING.glob("proposed-*.json")):
        data = json.loads(p.read_text(encoding="utf-8"))
        try:
            age = (datetime.now() - datetime.fromisoformat(data.get("created", ""))).days
        except ValueError:
            age = ""
        mt = data.get("meeting") or {}
        from_minutes = (mt.get("title", "") + (f" ({tracker.shown(mt['date'])})" if mt.get("date") else "")) \
            if mt and not mt.get("applied") and mt.get("minutes_decision", "pending") == "pending" else ""
        for item in data.get("items", []):
            if item.get("applied") or item.get("decision") != "pending":
                continue
            v = item.get("values") or {}
            if item.get("op") == "new":
                what = f"New {v.get('type') or item.get('table', '').rstrip('s')}"
                text = v.get("description") or v.get("decision") or v.get("summary") or v.get("name") or ""
                extra = [f"{k}: {tracker.shown(v[k]) if k in tracker.DATE_FIELDS else v[k]}" for k in ("owner", "due_date")
                         if v.get(k)]
                text = " · ".join([text, *extra] + ([review.detail(item)] if review.detail(item) else []))
            else:
                what = f"{item.get('op', '').title()} {item.get('table', '')}"
                text = ", ".join(f"{k} → {tracker.shown(c) if k in tracker.DATE_FIELDS and c else c}"
                                 for k, c in (item.get("changes") or {}).items())
            out.append([p.name, item.get("n"), review.project_of(item) or "", what,
                        item.get("id") or "", text, shown_flags(item), item.get("match", ""), age,
                        None, item.get("invalid_reason") or item.get("conflict_reason") or "", from_minutes])
        for block, field, label in (("meeting", "minutes_decision", "Minutes"), ("report", "report_decision", "Report"),
                                    ("file_changes", "file_changes_decision", "File change")):
            b = data.get(block)
            if b and not b.get("applied") and b.get(field, "pending") == "pending":
                out.append([p.name, block, ", ".join(b.get("projects", [])), label, "",
                            f"{b.get('title', '')} — {b.get('draft') or ''}", "", "", age, None,
                            "Edit in Word if needed, close it, then decide", ""])
    return out


def add_list(ws, col, n_rows, values):
    """Add a drop-down of `values` to column `col`, rows 2 to `n_rows`."""
    dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True,
                        showErrorMessage=True, errorTitle="creaite", error="Choose a value from the list.")
    dv.add(f"{col}2:{col}{n_rows}")
    ws.add_data_validation(dv)


def lock_column(ws, col, n_rows, label):
    """Refuse edits in column `col` (rows 2 to `n_rows`): `label` is set by creaite."""
    dv = DataValidation(type="custom", formula1="FALSE", showErrorMessage=True, errorTitle="creaite",
                        error=f"{label} is set by creaite and can't be edited here.")
    dv.add(f"{col}2:{col}{n_rows}")
    ws.add_data_validation(dv)


def rows_of(con, table, keys):
    """A table's rows as dicts, ordered by key. Labelled tables carry `labels` (the item's labels, comma-separated);
    a table an older database doesn't have yet (before `tracker.py migrate`) is empty."""
    try:
        rows = [dict(r) for r in con.execute(f"SELECT * FROM {table} ORDER BY {', '.join(keys)}")]
        if table in LABELLED:
            have = {}
            for r in con.execute("SELECT item_id, label FROM item_labels ORDER BY label COLLATE NOCASE"):
                have.setdefault(r["item_id"], []).append(r["label"])
            for r in rows:
                r["labels"] = ", ".join(have.get(r["id"], [])) or None
    except sqlite3.OperationalError:
        return []
    return rows


def cmd_export():
    """Write tracker-view.xlsx from tracker.db, one sheet per table plus a Review sheet, and record the export.

    Refuses (returns 3) while the workbook is open in Excel or has edits not yet imported.
    """
    if LOCK.exists():
        print("tracker-view.xlsx is open in Excel. Close it, then export again.")
        debuglog.warn("excel", "export paused: the workbook is open in Excel")
        return 3
    if has_unimported_edits():
        print("tracker-view.xlsx has edits that haven't been imported. Say \"import excel\" first; "
              "the workbook was not overwritten.")
        debuglog.warn("excel", "export paused: the workbook has edits that were not imported")
        return 3
    con = db()
    projects = [r["code"] for r in con.execute("SELECT code FROM projects ORDER BY code")]
    milestones = [r["id"] for r in con.execute("SELECT id FROM milestones ORDER BY id")]
    wb = Workbook()
    wb.remove(wb.active)
    snapshot = {}
    for name, table, keys, cols, locked in SHEETS:
        ws = wb.create_sheet(name)
        rows = rows_of(con, table, keys)
        snapshot[table] = {key_of(r, keys): r for r in rows}
        ws.append([label for label, _ in cols])
        for r in rows:
            ws.append([to_cell(f, r.get(f)) for _, f in cols])
        n_rows = len(rows) + 1 + EXTRA_ROWS
        for i, (label, field) in enumerate(cols, start=1):
            col = get_column_letter(i)
            head = ws.cell(row=1, column=i)
            head.font, head.fill = Font(bold=True), HEADER
            ws.column_dimensions[col].width = 48 if field in WIDE else 14 if field in DATES else 12
            if field in WIDE:
                for c in ws[f"{col}2:{col}{len(rows) + 1}"]:
                    c[0].alignment = Alignment(wrap_text=True, vertical="top")
            if field in DATES:
                for c in ws[f"{col}2:{col}{n_rows}"]:
                    c[0].number_format = "dd-mm-yyyy"
            if field in locked and not (table == "updates" and field in keys):
                lock_column(ws, col, n_rows, label)
                for c in ws[f"{col}2:{col}{len(rows) + 1}"]:
                    c[0].fill = LOCKED
            elif field in keys:  # updates: key cells of existing rows are locked; new rows may fill them
                for c in ws[f"{col}2:{col}{len(rows) + 1}"]:
                    c[0].fill = LOCKED
            if field == "project" and projects:
                add_list(ws, col, n_rows, projects)
            elif field == "milestone" and milestones and len(",".join(milestones)) < 250:
                add_list(ws, col, n_rows, milestones)   # optional link; Excel lists are limited to 255 characters
            elif field in YES_NO:
                add_list(ws, col, n_rows, ["Yes", "No"])
            elif choices(table, field):
                add_list(ws, col, n_rows, choices(table, field))
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(rows) + 1}"
        if table in READ_ONLY:
            ws.cell(row=1, column=len(cols) + 2, value="Read-only in Excel at first: links are made and removed in the "
                    "dashboard or with cx-edit. Edits on this sheet are ignored.").font = Font(italic=True)

    ws = wb.create_sheet("Review", 0)
    ws.append([c for c, _ in REVIEW_COLS])
    review = pending_rows()
    for r in review:
        ws.append(r)
    for i, (label, width) in enumerate(REVIEW_COLS, start=1):
        col = get_column_letter(i)
        ws.column_dimensions[col].width = width
        head = ws.cell(row=1, column=i)
        head.font, head.fill = Font(bold=True), HEADER
    if review:
        add_list(ws, "J", len(review) + 1, ["Approve", "Reject", "Edit"])
        for col in "ABCDEFGHIL":
            lock_column(ws, col, len(review) + 1, "This column")
    ws.freeze_panes = "A2"
    ws.cell(row=len(review) + 3, column=1,
            value="Decide here, save, then say \"import excel\" in Claude Code. Minutes, reports, new decisions "
                  "and new risks are approved one row at a time." if review else "Nothing is waiting for review.")

    tmp = paths.EXCEL_VIEW.with_name(paths.EXCEL_VIEW.stem + ".tmp.xlsx")
    wb.save(tmp)
    tmp.replace(paths.EXCEL_VIEW)
    save_state({"exported_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"), "hash": file_hash(),
                "imported_hash": None, "rows": snapshot})
    print(f"Exported {paths.rel(paths.EXCEL_VIEW)}: "
          + ", ".join(f"{len(snapshot[t])} {t}" for _, t, *_ in SHEETS) + f", {len(review)} to review.")
    debuglog.info("excel", "exported " + ", ".join(f"{t} {len(snapshot[t])}" for _, t, *_ in SHEETS) + f", {len(review)} to review")
    return 0


# ── Import ────────────────────────────────────────────────────────────────────

def read_sheet(ws, cols):
    """Rows as {field: stored value}; header labels map to fields, so column order may change."""
    header = [c.value for c in ws[1]]
    by_label = {label: field for label, field in cols}
    index = {i: by_label[h] for i, h in enumerate(header) if h in by_label}
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        rec = {index[i]: from_cell(index[i], v) for i, v in enumerate(row) if i in index}
        if any(v is not None for v in rec.values()):
            rows.append(rec)
    return rows


def new_proposal_path():
    """A free path for a new Excel-import proposal in system/staging/."""
    stamp = datetime.now().strftime("%Y-%m-%d-%H%M")
    p = paths.STAGING / f"proposed-{stamp}-excel.json"
    n = 2
    while p.exists():
        p = paths.STAGING / f"proposed-{stamp}-excel-{n}.json"
        n += 1
    return p


# cmd_import: stage the workbook's differences as a proposal (source_run excel-import) and copy Review sheet decisions into the
# pending proposals. Writes system/staging/ only; nothing reaches the database until apply. (A comment, not a docstring: the
# function is at its pinned length in dev/check-baseline.json.)
def cmd_import():
    if not paths.EXCEL_VIEW.exists():
        print("No tracker-view.xlsx to import.")
        return 0
    state = load_state()
    if not state.get("rows"):
        print("No export record found (system/staging/export-state.json). Nothing was imported.")
        debuglog.warn("excel", "import: no export record found, nothing imported")
        return 1
    if not has_unimported_edits():
        print("No unimported edits.")
        debuglog.info("excel", "import: no unimported edits")
        return 0
    if LOCK.exists():
        print("Note: the workbook is open in Excel; importing the last saved version.")
    wb = load_workbook(paths.EXCEL_VIEW, data_only=True)
    today = date.today().isoformat()
    items, notes = [], []

    for name, table, keys, cols, locked in SHEETS:
        if name not in wb.sheetnames:
            if table not in READ_ONLY and table in state["rows"]:
                notes.append(f"{name} sheet missing; skipped")
            continue
        if table in READ_ONLY:
            continue
        exported = state["rows"].get(table, {})
        seen = set()
        for rec in read_sheet(wb[name], cols):
            key = key_of(rec, keys) if all(rec.get(k) for k in keys) else None
            if table != "updates" and not rec.get(keys[0]):
                key = None
            if key and key in exported:
                seen.add(key)
                old = exported[key]
                changes = {f: rec.get(f) if rec.get(f) is not None else blank_default(table, f) for _, f in cols
                           if f not in locked and f not in keys and not same(rec.get(f), old.get(f), f)}
                if table == "milestones" and "baseline_date" in changes and old.get("baseline_date"):   # BL-186
                    changes.pop("baseline_date")
                    notes.append(f"{name}: {key}'s Baseline change was ignored: changing an agreed date is a rebaseline, which "
                                 "records a decision (the Milestones page's Rebaseline, or cx-edit)")
                if changes:
                    item = {"op": "update", "table": table, "id": key, "changes": changes}
                    if old.get("updated_at"):
                        item["based_on"] = old["updated_at"]
                    items.append(item)
            elif key and table != "updates":
                notes.append(f"{name}: {key} isn't in the last export; row ignored (IDs are set by creaite)")
            elif table == "projects":
                notes.append(f"{name}: new project rows are ignored; add projects with cx-edit (\"add project …\")")
            else:
                values = {f: v for f, v in rec.items() if v is not None and f not in ("id", "updated_at", "priority")}
                if table in ("actions", "raid", "promises"):
                    values.setdefault("raised", today)
                if table in ("actions", "raid", "decisions", "milestones", "workstreams", "promises"):
                    values["source"] = f"Recorded by Alan in Excel, {today}"
                items.append({"op": "new", "table": table, "values": values})
        for key in exported:
            if key not in seen:
                if table in ("actions", "raid"):
                    items.append({"op": "cancel", "table": table, "id": key, "confirm": "row deleted in Excel",
                                  "based_on": exported[key].get("updated_at")})
                else:
                    notes.append(f"{name}: {key} was deleted in Excel; {table} can't be deleted, "
                                 "so it returns at the next export")

    decided = apply_review_decisions(wb, notes)

    written = None
    if items:
        for n, item in enumerate(items, start=1):
            project = (item.get("values") or {}).get("project") or (item.get("id") or "").split("-")[0]
            flags = ["low-confidence"] if item["op"] == "cancel" else []
            item.update({"n": n, "match": "strong" if project else "unresolved", "flags": flags,
                         "decision": "pending"})
        written = new_proposal_path()
        written.write_text(json.dumps({"created": datetime.now().strftime("%Y-%m-%dT%H:%M:%S"),
                                       "source_run": "excel-import", "files": [paths.rel(paths.EXCEL_VIEW)],
                                       "items": items}, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
        # Flag rule failures and conflicts straight away
        subprocess.run([sys.executable, str(Path(__file__).with_name("tracker.py")), "validate", str(written)],
                       cwd=paths.ROOT, capture_output=True, text=True)

    state["imported_hash"] = file_hash()
    save_state(state)
    if written:
        data = json.loads(written.read_text(encoding="utf-8"))
        conflicts = sum("conflict" in i.get("flags", []) for i in data["items"])
        invalid = sum("invalid" in i.get("flags", []) for i in data["items"])
        counts = {op: sum(i["op"] == op for i in items) for op in ("new", "update", "cancel")}
        print(f"Staged {paths.rel(written)}: {counts['new']} new, {counts['update']} changed, "
              f"{counts['cancel']} deleted (to confirm), {conflicts} conflict(s), {invalid} invalid.")
        debuglog.info("excel", f"imported to {written.name}: {counts['new']} new, {counts['update']} changed, {counts['cancel']} deleted, "
                      f"{conflicts} conflict(s), {invalid} invalid, {len(notes)} note(s)")
    else:
        print("No table edits to stage.")
        debuglog.info("excel", f"import: no table edits to stage, {len(notes)} note(s)")
    if decided:
        print(f"Review sheet: {decided} decision(s) recorded in their proposals.")
        debuglog.info("excel", f"import: {decided} review decision(s) recorded")
    for n in notes:
        print("  " + n)
    return 0


def apply_review_decisions(wb, notes):
    """Copy Approve / Reject / Edit from the Review sheet into the pending proposal files."""
    if "Review" not in wb.sheetnames:
        return 0
    ws = wb["Review"]
    count = 0
    by_file = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0] or len(row) < 10 or not row[9]:
            continue
        by_file.setdefault(row[0], []).append((row[1], str(row[9]).strip().lower(), row[10] if len(row) > 10 else None))
    for fname, decisions in by_file.items():
        p = paths.STAGING / fname
        if not p.exists():
            notes.append(f"Review: {fname} is no longer pending; its decisions were ignored")
            continue
        data = json.loads(p.read_text(encoding="utf-8"))
        for n, choice, note in decisions:
            decision = {"approve": "approved", "reject": "rejected"}.get(choice)
            if n in ("meeting", "report", "file_changes"):
                block = data.get(n)
                field = {"meeting": "minutes_decision", "report": "report_decision"}.get(n, "file_changes_decision")
                if block and block.get(field, "pending") == "pending" and decision:
                    block[field], block["decided_in"] = decision, "excel"
                    count += 1
                continue
            item = next((i for i in data.get("items", []) if i.get("n") == n), None)
            if not item or item.get("decision") != "pending" or item.get("applied"):
                continue
            if decision:
                item["decision"], item["decided_in"] = decision, "excel"
                count += 1
            elif choice == "edit":
                item["review_note"] = f"Edit requested in Excel{': ' + note if note else ''}"
                count += 1
        p.write_text(json.dumps(data, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    return count


def cmd_changed():
    """Say whether the workbook has edits not yet imported; returns 1 if it has."""
    if has_unimported_edits():
        print("tracker-view.xlsx has unimported edits.")
        return 1
    print("No unimported edits.")
    return 0


def main(argv):
    """Run the export, import or changed command named in `argv`."""
    cmd = argv[0] if argv else ""
    if cmd == "export":
        return cmd_export()
    if cmd == "import":
        return cmd_import()
    if cmd == "changed":
        return cmd_changed()
    print(__doc__)
    return 2


if __name__ == "__main__":
    sys.exit(debuglog.main("excel_sync.py", main, sys.argv[1:]))
