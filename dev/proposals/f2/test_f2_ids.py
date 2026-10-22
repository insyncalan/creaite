"""F2 (IDs above 999): tests for the patched tree. Copy to dev/tests/ when the F2 diffs are applied; skipped on a tree without them.

Classes: next_id and the six CHECKs (a project past 999), A/B/C reader patterns (probe), D numeric SQL order, E JS comparator,
and the Excel export and import with a four-digit ID.
Run: python3 dev/tests/test_f2_ids.py
"""
import json
import re
import shutil
import sqlite3
import subprocess
import sys
import unittest
from pathlib import Path

HERE = Path(__file__).resolve().parent
TESTS = next(p for p in (HERE, HERE.parent, HERE.parent.parent, HERE.parent.parent / "tests") if (p / "test_tracker.py").exists())
sys.path.insert(0, str(TESTS))
from test_tracker import REPO, Vault, action  # noqa: E402

PATCHED = "instr(id, '-A-')" in (REPO / "system/scripts/schema.sql").read_text(encoding="utf-8")
SCRIPTS = REPO / "system/scripts"
SRC = REPO / "system/templates/dashboard/src"
SEARCH_ALLOWED = {"search.py"}     # its query is built from several tables and has no project column to lead with; left as it was


@unittest.skipUnless(PATCHED, "the F2 diffs are not applied in this tree")
class IdsAbove999(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.v = Vault()
        cls.v.run("init")
        cls.v.propose("big", [action(n) for n in range(1, 1003)])        # numbers 001 .. 1002 in one project
        cls.v.run("apply")

    @classmethod
    def tearDownClass(cls):
        cls.v.cleanup()

    def ids(self, sql):
        con = sqlite3.connect(self.v.root / "system/tracker.db")
        try:
            return [r[0] for r in con.execute(sql)]
        finally:
            con.close()

    def test_all_1002_written_and_next_is_1003(self):
        self.assertEqual(len(self.ids("SELECT id FROM actions")), 1002)
        self.assertIn("ABC-A-1000", self.ids("SELECT id FROM actions"))
        self.v.propose("one-more", [action(1, description="One more")])
        self.v.run("apply")
        self.assertIn("ABC-A-1003", self.ids("SELECT id FROM actions"))

    def test_a_five_digit_id_is_allowed_and_six_is_not(self):
        con = sqlite3.connect(self.v.root / "system/tracker.db")
        ins = "INSERT INTO actions (id, project, description, raised, source) VALUES (?, 'ABC', 'd', '2026-10-01', 's')"
        try:
            con.execute(ins, ("ABC-A-99999",))
            for bad in ("ABC-A-100000", "ABC-A-12", "ABC-A-12a4", "ABC-01234"):
                with self.assertRaises(sqlite3.IntegrityError, msg=bad):
                    con.execute(ins, (bad,))
        finally:
            con.close()

    def test_class_d_numeric_order(self):
        got = self.ids("SELECT id FROM actions ORDER BY project, length(id), id")
        nums = [int(i.rsplit("-", 1)[1]) for i in got]
        self.assertEqual(nums, sorted(nums))
        self.assertLess(got.index("ABC-A-999"), got.index("ABC-A-1000"))
        text = self.ids("SELECT id FROM actions ORDER BY id")
        self.assertGreater(text.index("ABC-A-999"), text.index("ABC-A-1000"))      # the old order is what the fix replaces

    def test_class_d_no_bare_order_by_id_left(self):
        left = []
        for f in sorted(SCRIPTS.glob("*.py")):
            if f.name in SEARCH_ALLOWED:
                continue
            for n, line in enumerate(f.read_text(encoding="utf-8").splitlines(), 1):
                if "ORDER BY" in line and re.search(r"ORDER BY [^\"]*\bid( DESC)?(?=\s*(LIMIT\b|\"|\)|$))", line) and "length(id)" not in line:
                    left.append(f"{f.name}:{n}")
        self.assertEqual(left, [])

    def test_class_e_js_comparator(self):
        probe = ("import { idCmp } from %s; console.log(JSON.stringify(['ABC-A-1000','ABC-A-999','ABC-A-1001','ABC-A-002','ABC-003']"
                 ".sort(idCmp)));" % json.dumps((SRC / "core/idsort.js").as_uri()))
        r = subprocess.run(["node", "--input-type=module", "-e", probe], capture_output=True, text=True)
        self.assertEqual(r.returncode, 0, r.stderr)
        got = json.loads(r.stdout)
        self.assertLess(got.index("ABC-A-999"), got.index("ABC-A-1000"))
        self.assertLess(got.index("ABC-A-1000"), got.index("ABC-A-1001"))

    def test_class_e_no_id_text_compare_left_in_the_dashboard(self):
        left = [f"{f.relative_to(SRC)}:{n}" for f in SRC.rglob("*") if f.suffix in (".js", ".jsx") and f.name != "idsort.js"
                for n, line in enumerate(f.read_text(encoding="utf-8").splitlines(), 1) if re.search(r"\ba(\[1\])?\.id\.localeCompare\(", line)]
        self.assertEqual(left, [])

    def test_class_abc_readers_see_four_digit_ids_whole(self):
        code = ("import sys, json; sys.path.insert(0, 'system/scripts'); t = 'See ABC-A-1000, ABC-1001 and ABC-D-1002.'; out = {}\n"
                "for mod, name in (('tracker','ITEM_ID'),('claims','ID'),('item_history','ITEM_RE'),('status_facts','ID_RE'),('prep_view','ID_RE'),"
                "('extraction','ITEM_ID'),('drafts','ID_RE'),('edit_ai','ID_RE'),('prep_shape','ID_RE'),('notes_find','ID_RE')):\n"
                "    out[mod+'.'+name] = [m.group(0) for m in getattr(__import__(mod), name).finditer(t)]\n"
                "print(json.dumps(out))")
        r = subprocess.run([sys.executable, "-c", code], cwd=self.v.root, capture_output=True, text=True, check=True)
        for name, found in json.loads(r.stdout).items():
            self.assertEqual(found, ["ABC-A-1000", "ABC-1001", "ABC-D-1002"], name)

    def test_excel_export_and_import_with_a_four_digit_id(self):
        sys.path.insert(0, str(self.v.root / "system/scripts/vendor"))
        from openpyxl import load_workbook
        self.v.run("export", script="excel_sync.py")
        path = self.v.root / "tracker-view.xlsx"
        wb = load_workbook(path)
        ws = next(wb[n] for n in wb.sheetnames if n.lower().startswith("action"))
        head = [c.value for c in ws[1]]
        col = {h: i + 1 for i, h in enumerate(head)}
        idcol = next(i for h, i in col.items() if str(h).lower() in ("id", "action id"))
        rows = {ws.cell(r, idcol).value: r for r in range(2, ws.max_row + 1) if ws.cell(r, idcol).value}
        self.assertIn("ABC-A-1000", rows)
        self.assertIn("ABC-A-0999" if False else "ABC-A-999", rows)
        self.assertGreaterEqual(len(rows), 1002)
        pcol = next(i for h, i in col.items() if str(h).lower() == "priority")
        ws.cell(rows["ABC-A-1000"], pcol).value = "Urgent"
        wb.save(path)
        out = self.v.run("import", script="excel_sync.py").stdout
        staged = [json.loads(p.read_text()) for p in (self.v.root / "system/staging").glob("proposed-*excel*.json")]
        items = [i for s in staged for i in s["items"]]
        self.assertTrue(any(i.get("id") == "ABC-A-1000" and i["changes"].get("priority") == "Urgent" for i in items), out)
        self.assertFalse(any(i.get("id") == "ABC-A-100" for i in items))


if __name__ == "__main__":
    unittest.main()
