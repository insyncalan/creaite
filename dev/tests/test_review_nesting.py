"""Review: items nested under their minutes (BL-41) — the dashboard, review.py list and the Excel Review sheet.

Run: python3 dev/tests/test_review_nesting.py   (the dashboard part needs Google Chrome; skipped otherwise)

The fictional portfolio from test_dashboard.py has one pending cx-minutes proposal (ABC Steering) holding a routine
action and a new risk for ABC, plus an item from Teams that needs a project.
"""
import json
import re
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from test_dashboard import CHROME, populate  # noqa: E402
from test_tracker import Vault  # noqa: E402

SCRIPT = r"""
<script>
(async()=>{
  const FILES=JSON.parse(document.getElementById("seed").textContent), wait=ms=>new Promise(r=>setTimeout(r,ms)),
    until=async(ok,ms=3000)=>{ for(const t0=Date.now(); !ok() && Date.now()-t0<ms;) await wait(50); }   /* slow CI runners: wait for the state, not a fixed time */, out={errors:[]};
  VW={ async read(p){ if(!(p in FILES)) throw new Error("not found: "+p); return FILES[p]; }, async write(p,t){ FILES[p]=t; },
       async remove(p){ delete FILES[p]; }, async exists(p){ return p in FILES; }, async append(p,t){ FILES[p]=(FILES[p]||"")+t; } };
  VSTATE="ok"; needVault=async()=>true;   /* the page's own vault check can finish late and reset the fake vault (slow CI) */ S.autoApply=false;
  out.snacks=[]; { const sn=snack; snack=function(m,u,l){ out.snacks.push(m+(u?" ["+(l||"Undo")+"]":"")); return sn(m,u,l); }; }   /* the record, if CI fails */
  const rowsNow=()=>[...document.querySelectorAll("#reviewBody tr[data-n]")].map(tr=>[tr.className.includes("mparent")?"parent":tr.className.includes("mchild")?"child":"row",tr.querySelector("td.title").textContent.replace(/\s+/g," ").trim().slice(0,40)]);
  try{
    go("review"); out.layout=rowsNow();
    const m=REV.find(r=>r.kind==="draft"&&r.ref==="meeting");
    out.toggle=document.querySelector("tr.mparent [data-act=t]").textContent; out.togOpen=document.querySelector("tr.mparent [data-act=t]").getAttribute("aria-expanded");
    document.querySelector("tr.mparent [data-act=t]").click(); out.collapsed=rowsNow().length;
    document.querySelector("tr.mparent [data-act=t]").click(); out.expanded=rowsNow().length;
    // ticking the parent ticks its routine children only
    const cb=document.querySelector("tr.mparent input[type=checkbox]"); cb.click();
    out.ticked=[...CHK].map(n=>REV.find(r=>r.n===n).rec.t); CHK.clear(); renderReview();
    // X on the parent does the same
    RSEL=[...document.querySelectorAll("#reviewBody tr[data-n]")].findIndex(tr=>tr.classList.contains("mparent"));
    document.dispatchEvent(new KeyboardEvent("keydown",{key:"x",bubbles:true})); out.xTicked=CHK.size; CHK.clear(); renderReview();
    // group approve: routine items only; the risk and the minutes stay
    document.querySelector("tr.mparent [data-act=ga]").click();
    await until(()=>REV.some(r=>r.file===m.file&&r.done==="approved")&&/approving one by one/.test(document.getElementById("snackT").textContent));
    out.afterGroup=REV.filter(r=>r.file===m.file).map(r=>[r.kind,r.done||"pending"]); out.groupSnack=document.getElementById("snackT").textContent;
    // rejecting the minutes offers to reject what's left
    document.querySelector("tr.mparent [data-act=r]").click();
    await until(()=>/Reject them too/.test(document.getElementById("snackU").textContent));
    out.rejectSnack=[document.getElementById("snackT").textContent,document.getElementById("snackU").textContent];
    await UNDO(); await until(()=>!REV.some(r=>r.file===m.file&&!r.done)); out.afterRejectAll=REV.filter(r=>r.file===m.file).map(r=>[r.kind,r.done||"pending"]);
  }catch(e){ out.errors.push(String(e&&e.stack||e)); }
  const pre=document.createElement("pre"); pre.id="result"; pre.textContent=JSON.stringify(out); document.body.appendChild(pre);
})();
</script>
"""


class ReviewNestingTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.v = v = Vault()
        populate(v)
        v.run(script="build_dashboard.py")

    @classmethod
    def tearDownClass(cls):
        cls.v.cleanup()

    def test_review_list_nests_items_under_minutes(self):
        out = self.v.run("list", script="review.py").stdout
        lines = [ln for ln in out.splitlines() if re.match(r"\s*(↳ )?\d+ ·", ln)]
        minutes = next(i for i, ln in enumerate(lines) if "Minutes" in ln)
        under = lines[minutes + 1] + lines[minutes + 2]                               # one-by-one items first
        self.assertTrue("Chase legal on the SOW" in under and "Vendor exits the market" in under)
        self.assertTrue(lines[minutes + 1].startswith("   ↳ ") and lines[minutes + 2].startswith("   ↳ "))
        self.assertNotIn("↳", next(ln for ln in lines if "Confirm data owner" in ln))   # needs a project: not nested
        numbers = [int(re.search(r"(\d+) ·", ln).group(1)) for ln in lines]
        self.assertEqual(numbers, sorted(numbers))                                      # still numbered in order

    def test_excel_review_sheet_from_minutes_column(self):
        sys.path.insert(0, str(self.v.root / "system" / "scripts"))
        import paths  # noqa: F401  (adds the bundled libraries)
        self.v.run("export", script="excel_sync.py")
        from openpyxl import load_workbook
        ws = load_workbook(self.v.root / "tracker-view.xlsx")["Review"]
        head = [c.value for c in ws[1]]
        self.assertEqual(head[9:12], ["Decision", "Note", "From minutes"])            # J and K unchanged
        by_text = {r[5]: r[11] for r in ws.iter_rows(min_row=2, values_only=True) if r[0]}
        sow = next(k for k in by_text if k.startswith("Chase legal on the SOW"))
        self.assertTrue(by_text[sow].startswith("ABC Steering"))
        owner = next(k for k in by_text if k.startswith("Confirm data owner"))
        self.assertEqual(by_text[owner], "ABC Steering")                              # same minutes, no project yet

    def test_excel_review_sheet_projects_and_dates(self):
        """Weekly-update items show their project (not "2026"); dates are DD-MM-YYYY."""
        sys.path.insert(0, str(self.v.root / "system" / "scripts"))
        import paths  # noqa: F401
        p = self.v.propose("xl", [{"n": 1, "op": "update", "table": "updates", "id": "2026-10-09|ABC", "match": "strong",
                                   "flags": [], "decision": "pending", "changes": {"summary": "Quiet week"}},
                                  {"n": 2, "op": "new", "table": "actions", "match": "strong", "flags": [],
                                   "decision": "pending", "values": {"project": "ABC", "description": "Dated action",
                                                                     "due_date": "2026-10-14", "raised": "2026-10-01",
                                                                     "source": "s"}}])
        try:
            self.v.run("export", script="excel_sync.py")
            from openpyxl import load_workbook
            ws = load_workbook(self.v.root / "tracker-view.xlsx")["Review"]
            rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] == p.name]
            self.assertEqual(rows[0][2], "ABC")
            self.assertIn("due_date: 14-10-2026", rows[1][5])
        finally:
            p.unlink()

    @unittest.skipUnless(CHROME.exists(), "Google Chrome not installed")
    def test_dashboard_nesting_and_group_actions(self):
        v = self.v
        seed = {f"system/staging/{p.name}": p.read_text(encoding="utf-8")
                for p in (v.root / "system/staging").glob("proposed-*.json")}
        html = (v.root / "dashboard.html").read_text(encoding="utf-8")
        tag = '<script type="application/json" id="seed">' + json.dumps(seed).replace("</", "<\\/") + "</script>"
        harness = v.root / "harness-nesting.html"
        harness.write_text(html.replace("</body>", tag + SCRIPT + "</body>"), encoding="utf-8")
        profile = tempfile.mkdtemp(prefix="creaite-chrome-")
        cmd = [str(CHROME), "--headless", "--disable-gpu", f"--user-data-dir={profile}", "--no-first-run",
               "--disable-component-update", "--disable-background-networking", "--virtual-time-budget=8000",
               "--dump-dom", harness.as_uri()]
        try:
            out = subprocess.run(cmd, capture_output=True, timeout=60).stdout
        except subprocess.TimeoutExpired as e:
            out = e.stdout or b""
        finally:
            shutil.rmtree(profile, ignore_errors=True)
        m = re.search(r'<pre id="result">(.*?)</pre>', out.decode("utf-8", "ignore"), re.S)
        self.assertTrue(m, "the page didn't finish")
        r = json.loads(m.group(1).replace("&lt;", "<").replace("&gt;", ">").replace("&quot;", '"').replace("&amp;", "&"))
        self.assertEqual(r["errors"], [])
        kinds = [k for k, _ in r["layout"]]
        p = kinds.index("parent")
        self.assertEqual(kinds[p:p + 3], ["parent", "child", "child"])               # both ABC items under the minutes
        self.assertTrue(any("Confirm data owner" in t for k, t in r["layout"] if k == "row"))   # not nested
        self.assertEqual((r["toggle"].strip(), r["togOpen"]), ("2 items", "true"))   # BL-158: the chevron is an icon, open is aria-expanded
        self.assertEqual(r["expanded"] - r["collapsed"], 2)                          # collapse hides the children
        self.assertEqual(r["ticked"], ["Chase legal on the SOW"])                    # the new risk is never ticked
        self.assertEqual(r["xTicked"], 1)
        self.assertEqual(sorted(r["afterGroup"]), [["actions", "approved"], ["actions", "pending"], ["draft", "pending"],
                                                   ["raid", "pending"]])                     # the DEF item needs a project
        self.assertIn("1 still need approving one by one", r["groupSnack"])
        self.assertEqual(r["rejectSnack"], ["Minutes rejected. Its 2 items stay for review.", "Reject them too"])
        self.assertEqual(sorted(r["afterRejectAll"]), [["actions", "approved"], ["actions", "rejected"], ["draft", "rejected"],
                                                       ["raid", "rejected"]], "messages shown: " + " | ".join(r["snacks"]))


if __name__ == "__main__":
    unittest.main(verbosity=2)
