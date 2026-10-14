"""Tests for milestones (BL-11): IDs, optional links from actions and RAID, Excel, migration.

Run: python3 dev/tests/test_milestones.py
"""
import json
import sqlite3
import sys
import unittest
from contextlib import closing
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from test_tracker import Vault, action  # noqa: E402

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "system" / "scripts" / "vendor"))
from openpyxl import load_workbook  # noqa: E402


def milestone(n, name, base, fc=None, **v):
    return {"n": n, "op": "new", "table": "milestones", "match": "strong", "flags": [], "decision": "approved",
            "values": {"project": "ABC", "name": name, "baseline_date": base, "forecast_date": fc or base,
                       "source": "Alan (cx-edit), 2026-10-05", **v}}


class MilestoneTests(unittest.TestCase):
    def setUp(self):
        self.v = Vault()
        self.v.run("init")
        self.v.propose("m", [milestone(1, "UAT complete", "2026-11-20"), milestone(2, "Go-live", "2026-12-11")])
        self.v.run("apply")

    def tearDown(self):
        self.v.cleanup()

    def test_ids_and_link_is_optional(self):
        self.assertEqual([m["id"] for m in self.v.rows("milestones")], ["ABC-M-001", "ABC-M-002"])
        self.v.propose("a", [action(1), action(2, description="Book UAT environment", milestone="ABC-M-001")])
        self.v.run("apply")
        rows = {a["id"]: a for a in self.v.rows("actions")}
        self.assertIsNone(rows["ABC-A-001"]["milestone"])          # no milestone: fine
        self.assertEqual(rows["ABC-A-002"]["milestone"], "ABC-M-001")
        p = self.v.propose("bad", [action(1, description="Ghost link", milestone="ABC-M-009")])
        r = self.v.run("validate", str(p), ok=False)
        self.assertIn("milestone ABC-M-009 is not in the tracker", r.stdout)

    def test_move_forecast_keeps_baseline_and_status_rules(self):
        self.v.propose("move", [{"n": 1, "op": "update", "table": "milestones", "id": "ABC-M-001",
                                 "changes": {"forecast_date": "2026-11-27"}, "decision": "approved"}])
        self.v.run("apply")
        m = self.v.rows("milestones")[0]
        self.assertEqual((m["baseline_date"], m["forecast_date"]), ("2026-11-20", "2026-11-27"))
        p = self.v.propose("st", [{"n": 1, "op": "update", "table": "milestones", "id": "ABC-M-001",
                                   "changes": {"status": "Late"}, "decision": "approved"}])
        self.assertIn("Planned, Done, Cancelled", self.v.run("validate", str(p), ok=False).stdout)   # BL-186: Missed is retired

    def test_excel_sheet_and_link_column(self):
        self.v.propose("a", [action(1, milestone="ABC-M-002")])
        self.v.run("apply")
        wb = load_workbook(self.v.root / "tracker-view.xlsx")
        self.assertIn("Milestones", wb.sheetnames)
        ms = wb["Milestones"]
        self.assertEqual([c.value for c in ms[1]][:7], ["ID", "Project", "Milestone", "Description", "Baseline", "Forecast", "Status"])
        acts = wb["Actions"]
        h = [c.value for c in acts[1]]
        self.assertEqual(acts.cell(row=2, column=h.index("Milestone") + 1).value, "ABC-M-002")
        # Unlinking in Excel stages a change to an empty milestone
        acts.cell(row=2, column=h.index("Milestone") + 1).value = None   # cell(value=None) wouldn't clear it
        wb.save(self.v.root / "tracker-view.xlsx")
        self.v.run("import", script="excel_sync.py")
        p = sorted((self.v.root / "system/staging").glob("proposed-*-excel*.json"))[-1]
        self.assertEqual(json.loads(p.read_text())["items"][0]["changes"], {"milestone": None})

    def test_description_applies_and_undoes(self):
        self.v.propose("d", [milestone(1, "Pilot start", "2026-12-18", description="Scope:\n- two sites\n- ten users")])
        self.v.run("apply")
        new = next(m for m in self.v.rows("milestones") if m["name"] == "Pilot start")
        self.assertEqual(new["description"], "Scope:\n- two sites\n- ten users")                  # newlines kept
        self.assertIsNone(self.v.rows("milestones")[0]["description"])                           # others stay empty
        self.v.propose("e", [{"n": 1, "op": "update", "table": "milestones", "id": "ABC-M-001",
                              "changes": {"description": "Sign-off by the steering committee"}, "decision": "approved"}])
        self.v.run("apply")
        self.assertEqual(self.v.rows("milestones")[0]["description"], "Sign-off by the steering committee")
        self.v.run("undo")
        self.assertIsNone(self.v.rows("milestones")[0]["description"])
        self.v.run("undo")
        self.assertNotIn("Pilot start", [m["name"] for m in self.v.rows("milestones")])

    def test_description_migration_keeps_every_row(self):
        db = self.v.root / "system/tracker.db"
        with closing(sqlite3.connect(db)) as con:                                  # a database from before the column
            con.execute("ALTER TABLE milestones DROP COLUMN description")
            con.commit()
            before = con.execute("SELECT id, project, name, baseline_date, forecast_date, status FROM milestones ORDER BY id").fetchall()
        self.assertIn("milestones", self.v.run(script="doctor.py", ok=False).stdout)         # doctor names it
        out = self.v.run("migrate").stdout
        self.assertIn("milestones (2 rows)", out)
        with closing(sqlite3.connect(db)) as con:
            cols = [r[1] for r in con.execute("PRAGMA table_info(milestones)")]
            self.assertEqual(cols[cols.index("name") + 1], "description")
            after = con.execute("SELECT id, project, name, baseline_date, forecast_date, status FROM milestones ORDER BY id").fetchall()
            self.assertIsNone(con.execute("SELECT description FROM milestones WHERE id='ABC-M-001'").fetchone()[0])
        self.assertEqual(before, after)
        self.assertIn("Schema is up to date", self.v.run("migrate").stdout)             # again: nothing changes

    def test_excel_round_trip_carries_the_description(self):
        self.v.propose("d", [{"n": 1, "op": "update", "table": "milestones", "id": "ABC-M-001",
                              "changes": {"description": "Line one\nLine two"}, "decision": "approved"}])
        self.v.run("apply")
        wb = load_workbook(self.v.root / "tracker-view.xlsx")
        ms = wb["Milestones"]
        h = [c.value for c in ms[1]]
        self.assertEqual(ms.cell(row=2, column=h.index("Description") + 1).value, "Line one\nLine two")
        ms.cell(row=2, column=h.index("Description") + 1).value = "Line one\nLine two\nLine three"
        wb.save(self.v.root / "tracker-view.xlsx")
        self.v.run("import", script="excel_sync.py")
        p = sorted((self.v.root / "system/staging").glob("proposed-*-excel*.json"))[-1]
        data = json.loads(p.read_text())
        self.assertEqual(data["items"][0]["changes"], {"description": "Line one\nLine two\nLine three"})
        data["items"][0]["decision"] = "approved"
        p.write_text(json.dumps(data))
        self.v.run("apply", str(p))
        self.assertEqual(self.v.rows("milestones")[0]["description"], "Line one\nLine two\nLine three")


if __name__ == "__main__":
    unittest.main(verbosity=2)
