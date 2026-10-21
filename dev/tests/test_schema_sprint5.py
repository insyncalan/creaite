"""Tests for Sprint 5 item 1: the schema migration (BL-53 decision fields, BL-50 promises, BL-90 part 2 links, BL-94 labels).

Run: python3 dev/tests/test_schema_sprint5.py
Each class builds a throwaway vault with fictional data and runs the real scripts there.
"""
import json
import re
import sqlite3
import sys
import unittest
from contextlib import closing
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from test_tracker import ABC, REPO, Vault, action  # noqa: E402
from test_workstreams import add_project, sql  # noqa: E402

SRC = f"{ABC}/minutes/2026-10-01 Steering.docx — Decisions"


def item(n, table, op="new", **kw):
    base = {"n": n, "op": op, "table": table, "match": "strong", "flags": [], "decision": "approved"}
    base.update(kw)
    return base


def decision(n, **v):
    return item(n, "decisions", values={"project": "ABC", "decision": "Extend the pilot by two weeks", "source": SRC, **v})


def promise(n, **v):
    return item(n, "promises", values={"project": "ABC", "promised_by": "Alan Carter", "promised_to": "Sara Example",
                                       "description": "Send the SOW", "due_date": "2026-10-09", "raised": "2026-10-05",
                                       "source": "Email: SOW, 2026-10-05", **v})


def raid(n, **v):
    return item(n, "raid", values={"project": "ABC", "type": "Dependency", "description": "Data feed from the client",
                                   "raised": "2026-10-01", "source": SRC, **v})


class Base(unittest.TestCase):
    rebuild = ()                    # views an apply may rebuild in the test vault (each costs seconds); the Excel tests keep excel_sync.py

    def setUp(self):
        self.v = Vault()
        for script in ("build_dashboard.py", "excel_sync.py"):
            if script not in self.rebuild:
                (self.v.root / "system/scripts" / script).unlink()          # after_apply skips a view whose script isn't there
        self.v.run("init")
        self.n = 0

    def tearDown(self):
        self.v.cleanup()

    def apply(self, *items, name=None):
        self.n += 1
        p = self.v.propose(name or f"t{self.n}", list(items))
        return self.v.run("apply", str(p))

    def refused(self, *items):
        """The schema refuses the item with a reason and writes nothing; returns the reason text.

        #280 (agreed in the review of v254-280-4): an invalid item no longer stops the whole apply. With tracker.py.diff it is
        held in Review with its reason and Apply exits 0 ("Held in Review, not applied: … <reason>"); before it, Apply stops
        (exit 1, "Apply stopped"). Either way the rule holds: no row of the refused item reaches the tracker."""
        self.n += 1
        before = self.written()
        p = self.v.propose(f"r{self.n}", list(items))
        r = self.v.run("apply", str(p), ok=False)
        held = p.exists() and any(i.get("decision") == "pending" and "invalid" in i.get("flags", [])
                                  for i in json.loads(p.read_text(encoding="utf-8")).get("items", []))
        p.unlink(missing_ok=True)
        if r.returncode == 0:
            self.assertIn("Held in Review, not applied", r.stdout)
            self.assertTrue(held, "a held item goes back to pending, flagged invalid")
        else:
            self.assertEqual(r.returncode, 1, r.stdout + r.stderr)
            self.assertIn("Apply stopped", r.stdout)
        self.assertEqual(self.written(), before, "a refused item writes nothing")
        return r.stdout

    def written(self):
        """What the tracker holds that an item could change: every change_log row and every decision, link and label."""
        return [sql(self.v, q) for q in ("SELECT * FROM change_log", "SELECT * FROM decisions", "SELECT * FROM links",
                                         "SELECT * FROM item_labels", "SELECT * FROM labels", "SELECT * FROM promises")]

    def undo(self):
        return self.v.run("undo").stdout

    def labels_of(self, item_id):
        return [r["label"] for r in sql(self.v, "SELECT label FROM item_labels WHERE item_id = ? ORDER BY label", item_id)]


# ── The migration ─────────────────────────────────────────────────────────────

OLD_DECISIONS = """CREATE TABLE IF NOT EXISTS decisions (
  id           TEXT PRIMARY KEY CHECK (id GLOB '[A-Z]*-D-[0-9][0-9][0-9]' AND length(id) BETWEEN 9 AND 13 AND substr(id, 1, length(id) - 6) NOT GLOB '*[^A-Z0-9]*'),
  project      TEXT NOT NULL REFERENCES projects(code),
  decision     TEXT NOT NULL,
  made_by      TEXT,
  client_visible INTEGER NOT NULL DEFAULT 0 CHECK (client_visible IN (0,1)),
  client_message TEXT NOT NULL DEFAULT '',
  date         TEXT NOT NULL,
  source       TEXT NOT NULL,
  updated_at   TEXT NOT NULL DEFAULT (datetime('now'))
);

"""


def before_sprint_5_schema():
    """Today's schema with the decisions table as it was and none of the new tables (built from the checked-out file)."""
    s = (REPO / "system/scripts/schema.sql").read_text(encoding="utf-8")
    s = re.sub(r"-- Decisions \(BL-53\).*?\n\);\n\n", OLD_DECISIONS, s, flags=re.S)
    s = re.sub(r"-- Promises \(BL-50\).*?(?=CREATE TABLE IF NOT EXISTS updates)", "", s, flags=re.S)
    return s


class MigrationTests(Base):
    def setUp(self):
        self.v = v = Vault()
        old = before_sprint_5_schema()
        self.assertNotIn("promises", old)
        self.assertNotIn("needed_by", old)
        self.n = 0
        (v.root / "system/scripts/schema.sql").write_text(old)
        v.run("init")
        (v.root / "system/scripts/schema.sql").write_text((REPO / "system/scripts/schema.sql").read_text())
        with closing(sqlite3.connect(v.root / "system/tracker.db")) as con:
            for i, (d, by, cv) in enumerate([("2026-09-15", "Sara Example", 1), ("2026-10-01", None, 0)], start=1):
                con.execute("INSERT INTO decisions (id, project, decision, made_by, client_visible, client_message, date, source, "
                            "updated_at) VALUES (?, 'ABC', ?, ?, ?, ?, ?, 'test', '2026-10-02 09:00:00.000')",
                            (f"ABC-D-{i:03d}", f"Decision {i}", by, cv, "Go-live moves" if cv else "", d))
            con.execute("INSERT INTO actions (id, project, description, raised, source) VALUES "
                        "('ABC-A-001', 'ABC', 'An action', '2026-10-01', 'test')")
            con.commit()

    def snapshot(self):
        return {t: sql(self.v, f"SELECT * FROM {t} ORDER BY 1") for t in ("projects", "actions", "raid", "milestones")}

    def test_existing_decisions_become_made_and_keep_their_dates(self):
        before = self.snapshot()
        old = sql(self.v, "SELECT * FROM decisions ORDER BY id")
        out = self.v.run("migrate").stdout
        self.assertIn("Migrated decisions (2 rows)", out)
        now = sql(self.v, "SELECT * FROM decisions ORDER BY id")
        self.assertEqual([(d["id"], d["status"], d["date"]) for d in now],
                         [("ABC-D-001", "Made", "2026-09-15"), ("ABC-D-002", "Made", "2026-10-01")])
        for o, n in zip(old, now):               # every old column unchanged, the new ones empty
            self.assertEqual({k: n[k] for k in o}, o)
            self.assertEqual((n["needed_by"], n["reasoning"], n["options"], n["revisit"], n["replaces"]), (None,) * 5)
        self.assertEqual(self.snapshot(), before)

    def test_new_tables_and_the_sensitive_label(self):
        self.v.run("migrate")
        for t in ("promises", "links", "labels", "item_labels"):
            self.assertEqual(sql(self.v, f"SELECT count(*) AS n FROM {t}")[0]["n"] if t != "labels" else 1, 0 if t != "labels" else 1)
        self.assertEqual(sql(self.v, "SELECT * FROM labels"),
                         [{"name": "Sensitive", "colour": "red", "note": "Commercially or client sensitive",
                           "behaviour": "warn", "default_choice": "exclude", "builtin": 1}])

    def test_running_it_twice_changes_nothing(self):
        self.v.run("migrate")
        once = {t: sql(self.v, f"SELECT * FROM {t} ORDER BY 1") for t in ("decisions", "labels", "actions", "projects")}
        schema = sql(self.v, "SELECT name, sql FROM sqlite_master ORDER BY name")
        backups = len(list((self.v.root / "system/backups").glob("*-migrate.db")))
        out = self.v.run("migrate").stdout
        self.assertIn("Schema is up to date", out)
        self.assertEqual({t: sql(self.v, f"SELECT * FROM {t} ORDER BY 1") for t in once}, once)
        self.assertEqual(sql(self.v, "SELECT name, sql FROM sqlite_master ORDER BY name"), schema)
        self.assertEqual(len(list((self.v.root / "system/backups").glob("*-migrate.db"))), backups)   # no second rebuild

    def test_a_decision_can_be_pending_after_migrating_and_a_made_one_needs_its_date(self):
        self.v.run("migrate")
        self.apply(decision(1, status="Pending", made_by="Sara Example", needed_by="2026-10-16"))
        self.assertEqual(sql(self.v, "SELECT id, status, date FROM decisions WHERE id = 'ABC-D-003'"),
                         [{"id": "ABC-D-003", "status": "Pending", "date": None}])
        with closing(sqlite3.connect(self.v.root / "system/tracker.db")) as con:     # the schema itself refuses it too
            with self.assertRaises(sqlite3.IntegrityError):
                con.execute("INSERT INTO decisions (id, project, decision, source) VALUES ('ABC-D-009', 'ABC', 'x', 'y')")


# ── Decisions (BL-53) ─────────────────────────────────────────────────────────

class DecisionTests(Base):
    def test_pending_decision_then_made_keeps_its_number(self):
        self.apply(decision(1, status="Pending", made_by="Sara Example", needed_by="2026-10-16", reasoning="Pilot ends 09-10",
                            options="Extend; stop; go live early", revisit="revisit if adoption falls below 60%"))
        d = sql(self.v, "SELECT * FROM decisions")[0]
        self.assertEqual((d["id"], d["status"], d["date"], d["needed_by"], d["made_by"]),
                         ("ABC-D-001", "Pending", None, "2026-10-16", "Sara Example"))
        self.assertEqual((d["reasoning"], d["options"], d["revisit"]),
                         ("Pilot ends 09-10", "Extend; stop; go live early", "revisit if adoption falls below 60%"))
        self.apply(item(1, "decisions", "update", id="ABC-D-001", changes={"status": "Made", "date": "2026-10-12"}))
        d = sql(self.v, "SELECT * FROM decisions")[0]
        self.assertEqual((d["id"], d["status"], d["date"]), ("ABC-D-001", "Made", "2026-10-12"))
        self.assertEqual(self.undo().splitlines()[0], "Undid apply 2: 1 item(s) reversed.")      # back to Pending, no date
        d = sql(self.v, "SELECT * FROM decisions")[0]
        self.assertEqual((d["status"], d["date"]), ("Pending", None))

    def test_withdraw(self):
        self.apply(decision(1, status="Pending"))
        self.apply(item(1, "decisions", "update", id="ABC-D-001", changes={"status": "Withdrawn"}))
        self.assertEqual(sql(self.v, "SELECT status, date FROM decisions"), [{"status": "Withdrawn", "date": None}])

    def test_a_made_decision_without_a_date_is_refused(self):
        self.assertIn("needs its date", self.refused(decision(1)))                         # status defaults to Made
        self.assertIn("needs its date", self.refused(decision(1, status="Made")))
        self.apply(decision(1, status="Pending"))
        self.assertIn("needs its date", self.refused(item(1, "decisions", "update", id="ABC-D-001", changes={"status": "Made"})))
        self.assertEqual(sql(self.v, "SELECT status FROM decisions"), [{"status": "Pending"}])

    def test_a_pending_decision_has_no_date_and_a_made_one_cannot_go_back(self):
        self.assertIn("no date until it is made", self.refused(decision(1, status="Pending", date="2026-10-12")))
        self.apply(decision(1, date="2026-10-01"))
        self.assertIn("can't go back to Pending",
                      self.refused(item(1, "decisions", "update", id="ABC-D-001", changes={"status": "Pending", "date": None})))

    def test_bad_values_are_refused(self):
        self.assertIn("status must be one of", self.refused(decision(1, status="Maybe")))
        self.assertIn("needed_by must be a date", self.refused(decision(1, status="Pending", needed_by="16-10-2026")))

    def test_replaces(self):
        self.apply(decision(1, date="2026-09-01", decision="Pilot ends 09-10"))
        self.apply(decision(1, date="2026-10-12", decision="Pilot extended", replaces="ABC-D-001"))
        self.assertEqual(sql(self.v, "SELECT id, replaces FROM decisions ORDER BY id"),
                         [{"id": "ABC-D-001", "replaces": None}, {"id": "ABC-D-002", "replaces": "ABC-D-001"}])
        self.assertIn("is not in the tracker", self.refused(decision(1, date="2026-10-12", replaces="ABC-D-077")))
        self.assertIn("would replace each other",
                      self.refused(item(1, "decisions", "update", id="ABC-D-001", changes={"replaces": "ABC-D-002"})))

    def test_replaces_stays_in_the_project(self):
        add_project(self.v, "DEF", "Digital Enablement")
        self.v.run("init")
        self.apply(decision(1, project="DEF", date="2026-09-01"))
        self.assertIn("belongs to DEF, not ABC", self.refused(decision(1, date="2026-10-12", replaces="DEF-D-001")))


# ── Promises (BL-50) ──────────────────────────────────────────────────────────

class PromiseTests(Base):
    def test_new_done_dropped_and_undo(self):
        self.apply(promise(1), promise(2, description="Review the plan"))
        rows = sql(self.v, "SELECT id, status, closed, promised_by, promised_to FROM promises ORDER BY id")
        self.assertEqual([r["id"] for r in rows], ["ABC-P-001", "ABC-P-002"])            # house-style IDs, per project
        self.assertEqual((rows[0]["status"], rows[0]["closed"], rows[0]["promised_to"]), ("Open", None, "Sara Example"))
        self.apply(item(1, "promises", "update", id="ABC-P-001", changes={"status": "Done"}))
        done = sql(self.v, "SELECT status, closed FROM promises WHERE id = 'ABC-P-001'")[0]
        self.assertEqual(done["status"], "Done")
        self.assertTrue(done["closed"])
        self.apply(item(1, "promises", "cancel", id="ABC-P-002", confirm="no longer needed"))
        self.assertEqual(sql(self.v, "SELECT status FROM promises WHERE id = 'ABC-P-002'"), [{"status": "Dropped"}])
        self.undo()
        self.assertEqual(sql(self.v, "SELECT status, closed FROM promises WHERE id = 'ABC-P-002'"), [{"status": "Open", "closed": None}])
        self.undo()
        self.assertEqual(sql(self.v, "SELECT status, closed FROM promises WHERE id = 'ABC-P-001'"), [{"status": "Open", "closed": None}])
        self.undo()
        self.assertEqual(sql(self.v, "SELECT count(*) AS n FROM promises"), [{"n": 0}])

    def test_refusals(self):
        self.assertIn("missing required field(s): promised_to", self.refused(promise(1, promised_to="")))
        self.assertIn("status must be one of", self.refused(promise(1, status="Closed")))
        self.assertIn("project XYZ is not in the tracker", self.refused(promise(1, project="XYZ")))

    def test_no_priority_or_workstream(self):
        self.assertIn("unknown field(s) for promises: priority", self.refused(promise(1, priority="High")))
        self.assertIn("unknown field(s) for promises: workstream", self.refused(promise(1, workstream="Data")))


# ── Links (BL-90 part 2) ──────────────────────────────────────────────────────

def link(n, a, b, **v):
    return item(n, "links", values={"from_item": a, "to_item": b, "source": "Alan (cx-edit)", **v})


class LinkTests(Base):
    def setUp(self):
        super().setUp()
        add_project(self.v, "DEF", "Digital Enablement")
        self.v.run("init")
        self.apply(raid(1), action(2), action(3, project="DEF"))        # ABC-001 (a dependency), ABC-A-002, DEF-A-001

    def test_link_across_projects_and_undo(self):
        self.apply(link(1, "ABC-001", "ABC-A-002"), link(2, "ABC-001", "DEF-A-001"))
        rows = sql(self.v, "SELECT from_item, to_item, kind FROM links ORDER BY to_item")
        self.assertEqual(rows, [{"from_item": "ABC-001", "to_item": "ABC-A-002", "kind": "blocks"},
                                {"from_item": "ABC-001", "to_item": "DEF-A-001", "kind": "blocks"}])
        self.undo()
        self.assertEqual(sql(self.v, "SELECT count(*) AS n FROM links"), [{"n": 0}])

    def test_remove_a_link_and_undo_puts_it_back(self):
        self.apply(link(1, "ABC-001", "ABC-A-002"))
        before = sql(self.v, "SELECT * FROM links")
        self.apply(item(1, "links", "cancel", id="ABC-001|ABC-A-002|blocks", confirm="not needed"))
        self.assertEqual(sql(self.v, "SELECT count(*) AS n FROM links"), [{"n": 0}])
        self.undo()
        self.assertEqual(sql(self.v, "SELECT * FROM links"), before)

    def test_refusals(self):
        self.assertIn("ABC-099 is not in the tracker", self.refused(link(1, "ABC-099", "ABC-A-002")))
        self.assertIn("ABC-A-077 is not in the tracker", self.refused(link(1, "ABC-001", "ABC-A-077")))
        self.assertIn("isn't an action", self.refused(link(1, "ABC-001", "nonsense")))
        self.assertIn("can't block itself", self.refused(link(1, "ABC-001", "ABC-001")))
        self.assertIn("kind must be one of", self.refused(link(1, "ABC-001", "ABC-A-002", kind="relates")))
        self.apply(link(1, "ABC-001", "ABC-A-002"))
        self.assertIn("already has", self.refused(link(1, "ABC-001", "ABC-A-002")))

    def test_links_that_would_loop_are_refused(self):
        self.apply(link(1, "ABC-001", "ABC-A-002"), link(2, "ABC-A-002", "DEF-A-001"))
        self.assertIn("would loop", self.refused(link(1, "DEF-A-001", "ABC-001")))
        self.assertIn("would loop", self.refused(link(1, "ABC-A-002", "ABC-001")))

    def test_a_milestone_can_be_blocked(self):
        self.apply(item(1, "milestones", values={"project": "ABC", "name": "Go-live", "forecast_date": "2026-11-20",
                                                 "source": "Alan"}))
        self.apply(link(1, "ABC-001", "ABC-M-001"))
        self.assertEqual(sql(self.v, "SELECT to_item FROM links"), [{"to_item": "ABC-M-001"}])


# ── Labels (BL-94) ────────────────────────────────────────────────────────────

class LabelTests(Base):
    def test_sensitive_is_built_in_and_cannot_be_deleted_or_renamed(self):
        self.assertEqual(sql(self.v, "SELECT name, builtin FROM labels"), [{"name": "Sensitive", "builtin": 1}])
        self.assertIn("built-in label and can't be deleted",
                      self.refused(item(1, "labels", "cancel", id="Sensitive", confirm="x")))
        with closing(sqlite3.connect(self.v.root / "system/tracker.db")) as con:                 # the schema holds it too
            for statement in ("DELETE FROM labels WHERE name = 'Sensitive'",
                              "UPDATE labels SET name = 'Secret' WHERE name = 'Sensitive'",
                              "UPDATE labels SET builtin = 0 WHERE name = 'Sensitive'"):
                with self.assertRaises(sqlite3.DatabaseError, msg=statement):
                    con.execute(statement)
        self.assertIn("builtin", self.refused(item(1, "labels", "update", id="Sensitive", changes={"builtin": 0})))

    def test_new_label_and_delete_and_undo(self):
        self.apply(item(1, "labels", values={"name": "Legal review", "colour": "#7a3", "note": "Legal reads it first",
                                             "behaviour": "exclude"}))
        self.assertEqual(sql(self.v, "SELECT name, behaviour, default_choice, builtin FROM labels WHERE name = 'Legal review'"),
                         [{"name": "Legal review", "behaviour": "exclude", "default_choice": "include", "builtin": 0}])
        self.apply(item(1, "labels", "cancel", id="Legal review", confirm="x"))
        self.assertEqual(len(sql(self.v, "SELECT * FROM labels")), 1)
        self.undo()
        self.assertEqual(sql(self.v, "SELECT note FROM labels WHERE name = 'Legal review'"), [{"note": "Legal reads it first"}])

    def test_label_refusals(self):
        self.assertIn("built-in labels are created by creaite", self.refused(item(1, "labels", values={"name": "X", "builtin": 1})))
        self.assertIn("behaviour must be one of", self.refused(item(1, "labels", values={"name": "X", "behaviour": "shout"})))
        self.assertIn("comma", self.refused(item(1, "labels", values={"name": "A, B"})))
        self.assertIn("colour", self.refused(item(1, "labels", values={"name": "X", "colour": "red;}<script>"})))
        self.apply(item(1, "labels", values={"name": "Legal"}))
        self.assertIn("already has", self.refused(item(1, "labels", values={"name": "legal"})))             # names ignore case

    def test_a_label_in_use_cannot_be_deleted(self):
        self.apply(item(1, "labels", values={"name": "Legal"}), action(2, labels=["Legal"]))
        self.assertIn("1 item(s) still carry Legal", self.refused(item(1, "labels", "cancel", id="Legal", confirm="x")))

    def test_labels_on_new_items_and_on_edits(self):
        self.apply(action(1, labels=["sensitive"], label_reason="reads as criticism of the sponsor", label_suggested=1),
                   raid(2, labels="Sensitive"), decision(3, date="2026-10-01", labels=["Sensitive"]))
        self.assertEqual(self.labels_of("ABC-A-001"), ["Sensitive"])                       # stored in the label's own spelling
        self.assertEqual(self.labels_of("ABC-002"), ["Sensitive"])
        self.assertEqual(self.labels_of("ABC-D-003"), ["Sensitive"])
        row = sql(self.v, "SELECT * FROM item_labels WHERE item_id = 'ABC-A-001'")[0]
        self.assertEqual((row["reason"], row["suggested"]), ("reads as criticism of the sponsor", 1))
        self.apply(item(1, "labels", values={"name": "Legal"}))
        self.apply(item(1, "actions", "update", id="ABC-A-001", changes={"labels": "Sensitive, Legal"}))
        self.assertEqual(self.labels_of("ABC-A-001"), ["Legal", "Sensitive"])
        self.apply(item(1, "actions", "update", id="ABC-A-001", changes={"labels": ["Legal"]}))      # the list is the full set
        self.assertEqual(self.labels_of("ABC-A-001"), ["Legal"])
        self.apply(item(1, "actions", "update", id="ABC-A-001", changes={"labels": None, "latest_update": "Chased"}))
        self.assertEqual(self.labels_of("ABC-A-001"), [])
        self.assertEqual(sql(self.v, "SELECT latest_update FROM actions")[0]["latest_update"], "Chased")

    def test_label_changes_undo_with_the_apply(self):
        self.apply(action(1))
        self.apply(item(1, "actions", "update", id="ABC-A-001", changes={"labels": ["Sensitive"]}))
        self.assertEqual(self.labels_of("ABC-A-001"), ["Sensitive"])
        self.undo()
        self.assertEqual(self.labels_of("ABC-A-001"), [])
        self.apply(item(1, "actions", "update", id="ABC-A-001", changes={"labels": ["Sensitive"]}))
        self.apply(item(1, "actions", "update", id="ABC-A-001", changes={"labels": []}))
        self.assertEqual(self.labels_of("ABC-A-001"), [])
        self.undo()                                                                          # the removal comes back
        self.assertEqual(self.labels_of("ABC-A-001"), ["Sensitive"])
        self.undo()
        self.assertEqual(self.labels_of("ABC-A-001"), [])

    def test_a_new_item_with_labels_undoes_whole(self):
        self.apply(action(1, labels=["Sensitive"]))
        self.undo()
        self.assertEqual(sql(self.v, "SELECT count(*) AS n FROM actions"), [{"n": 0}])
        self.assertEqual(sql(self.v, "SELECT count(*) AS n FROM item_labels"), [{"n": 0}])

    def test_item_labels_rows_directly(self):
        self.apply(action(1))
        self.apply(item(1, "item_labels", values={"item_id": "ABC-A-001", "label": "sensitive", "reason": "pricing"}))
        self.assertEqual(sql(self.v, "SELECT label, reason, added_by, suggested FROM item_labels"),
                         [{"label": "Sensitive", "reason": "pricing", "added_by": "cx-edit", "suggested": 0}])
        self.apply(item(1, "item_labels", "cancel", id="ABC-A-001|Sensitive", confirm="x"))
        self.assertEqual(self.labels_of("ABC-A-001"), [])
        self.undo()
        self.assertEqual(self.labels_of("ABC-A-001"), ["Sensitive"])

    def test_refusals(self):
        self.assertIn("unknown label 'Secret'", self.refused(action(1, labels=["Secret"])))
        self.assertIn("unknown label 'Secret'", self.refused(raid(1, labels="Sensitive, Secret")))
        self.apply(action(1))
        self.assertIn("unknown label", self.refused(item(1, "actions", "update", id="ABC-A-001", changes={"labels": ["Nope"]})))
        self.assertIn("unknown label", self.refused(item(1, "item_labels", values={"item_id": "ABC-A-001", "label": "Nope"})))
        self.assertIn("isn't an action, RAID item or decision",
                      self.refused(item(1, "item_labels", values={"item_id": "ABC-A-099", "label": "Sensitive"})))
        self.assertIn("labels can only be put on actions, RAID items and decisions",
                      self.refused(promise(1, labels=["Sensitive"])))
        self.assertEqual(sql(self.v, "SELECT count(*) AS n FROM item_labels"), [{"n": 0}])   # nothing was half-written

    def test_a_rejected_link_or_label_is_not_proposed_again(self):
        self.apply(action(1), raid(2), raid(3, description="Another dependency"))
        a = link(1, "ABC-002", "ABC-A-001")
        b = link(2, "ABC-003", "ABC-A-001")
        a["decision"] = b["decision"] = "rejected"
        self.v.propose("rej", [a, b])
        self.v.run("apply")
        self.assertEqual(len(sql(self.v, "SELECT * FROM rejected")), 2)                     # each link has its own fingerprint


# ── Excel ─────────────────────────────────────────────────────────────────────

try:
    from openpyxl import load_workbook
except ImportError:                                      # the bundled copy
    sys.path.insert(0, str(REPO / "system/scripts/vendor"))
    from openpyxl import load_workbook


class ExcelTests(Base):
    rebuild = ("excel_sync.py",)

    def setUp(self):
        super().setUp()
        self.apply(action(1, labels=["Sensitive"]), raid(2),
                   decision(3, date="2026-10-01", reasoning="Pilot ran well", options="A; B", revisit="2026-12-01",
                            labels=["Sensitive"]),
                   decision(4, status="Pending", made_by="Sara Example", needed_by="2026-10-16"),
                   promise(5), link(6, "ABC-002", "ABC-A-001"))
        self.v.run("export", script="excel_sync.py")
        self.xlsx = self.v.root / "tracker-view.xlsx"

    def sheet(self, name):
        ws = load_workbook(self.xlsx)[name]
        head = [c.value for c in ws[1]]
        return head, [dict(zip(head, [c.value for c in row])) for row in ws.iter_rows(min_row=2) if row[0].value]

    def edit(self, name, change):
        wb = load_workbook(self.xlsx)
        ws = wb[name]
        head = [c.value for c in ws[1]]
        change(ws, head)
        wb.save(self.xlsx)

    def test_export_carries_the_new_fields_and_sheets(self):
        head, rows = self.sheet("Decisions")
        self.assertEqual(head[:12], ["ID", "Project", "Decision", "Status", "Made by", "Needed by", "Date", "Reasoning",
                                     "Options", "Revisit", "Replaces", "Client-visible"])
        self.assertIn("Labels", head)
        made, pending = rows
        self.assertEqual((made["Status"], made["Reasoning"], made["Options"], made["Revisit"], made["Labels"]),
                         ("Made", "Pilot ran well", "A; B", "2026-12-01", "Sensitive"))
        self.assertEqual((pending["Status"], pending["Made by"], pending["Date"], pending["Needed by"].date().isoformat()),
                         ("Pending", "Sara Example", None, "2026-10-16"))
        _, promises = self.sheet("Promises")
        self.assertEqual([(p["ID"], p["Promised by"], p["Promised to"], p["Status"]) for p in promises],
                         [("ABC-P-005", "Alan Carter", "Sara Example", "Open")])
        for name in ("Actions", "RAID"):
            head, rows = self.sheet(name)
            self.assertIn("Labels", head)
        self.assertEqual(self.sheet("Actions")[1][0]["Labels"], "Sensitive")
        self.assertIsNone(self.sheet("RAID")[1][0]["Labels"])

    def test_links_sheet_is_read_only_and_says_so(self):
        wb = load_workbook(self.xlsx)
        ws = wb["Links"]
        self.assertEqual([c.value for c in ws[1]][:5], ["Blocker", "Blocks", "Kind", "Created", "Source"])
        self.assertEqual([c.value for c in ws[2]][:3], ["ABC-002", "ABC-A-001", "blocks"])
        self.assertIn("Read-only in Excel", ws.cell(row=1, column=7).value)
        self.assertTrue(any(dv.type == "custom" for dv in ws.data_validations.dataValidation))      # cells locked as elsewhere
        self.edit("Links", lambda ws, head: ws.cell(row=2, column=2, value="ABC-A-099"))
        out = self.v.run("import", script="excel_sync.py").stdout
        self.assertIn("No table edits to stage", out)
        self.assertEqual(sql(self.v, "SELECT to_item FROM links"), [{"to_item": "ABC-A-001"}])

    def approve_and_apply(self):
        """Approve every item of the staged Excel proposal (Alan's click in Review), then apply it."""
        prop = next((self.v.root / "system/staging").glob("proposed-*-excel.json"))
        data = json.loads(prop.read_text())
        for i in data["items"]:
            i["decision"] = "approved"
        prop.write_text(json.dumps(data))
        return self.v.run("apply", str(prop))

    def test_import_edits_decision_fields_and_labels(self):
        def change(ws, head):
            for row in ws.iter_rows(min_row=2):
                if row[0].value == "ABC-D-004":
                    row[head.index("Status")].value = "Made"
                    row[head.index("Date")].value = "12-10-2026"
                    row[head.index("Reasoning")].value = "Sara confirmed on the call"
                    row[head.index("Replaces")].value = "ABC-D-003"
                if row[0].value == "ABC-D-003":
                    row[head.index("Labels")].value = None
        self.edit("Decisions", change)
        out = self.v.run("import", script="excel_sync.py").stdout
        self.assertIn("2 changed", out)
        self.approve_and_apply()
        d = sql(self.v, "SELECT * FROM decisions WHERE id = 'ABC-D-004'")[0]
        self.assertEqual((d["status"], d["date"], d["reasoning"], d["replaces"]),
                         ("Made", "2026-10-12", "Sara confirmed on the call", "ABC-D-003"))
        self.assertEqual(self.labels_of("ABC-D-003"), [])

    def test_labels_compare_as_a_set(self):
        self.apply(item(1, "labels", values={"name": "Legal"}), item(2, "actions", "update", id="ABC-A-001",
                                                                       changes={"labels": ["Sensitive", "Legal"]}))
        self.v.run("export", script="excel_sync.py")
        self.assertEqual(self.sheet("Actions")[1][0]["Labels"], "Legal, Sensitive")
        self.edit("Actions", lambda ws, head: setattr(ws.cell(row=2, column=head.index("Labels") + 1), "value", "sensitive,legal"))
        out = self.v.run("import", script="excel_sync.py").stdout
        self.assertIn("No table edits to stage", out)                                       # same labels, other order and case

    def test_an_unknown_label_is_flagged_invalid(self):
        self.edit("Actions", lambda ws, head: setattr(ws.cell(row=2, column=head.index("Labels") + 1), "value", "Sensitive, Secret"))
        self.v.run("import", script="excel_sync.py")
        prop = next((self.v.root / "system/staging").glob("proposed-*-excel.json"))
        got = json.loads(prop.read_text())["items"][0]
        self.assertIn("invalid", got["flags"])
        self.assertIn("unknown label 'Secret'", got["invalid_reason"])

    def test_a_new_made_decision_row_without_a_date_is_flagged_invalid(self):
        def change(ws, head):
            row = ws.max_row + 1
            for col, value in (("Project", "ABC"), ("Decision", "Switch to weekly releases"), ("Status", "Made")):
                ws.cell(row=row, column=head.index(col) + 1, value=value)
            ws.cell(row=row, column=1, value=None)
        self.edit("Decisions", change)
        self.v.run("import", script="excel_sync.py")
        prop = next((self.v.root / "system/staging").glob("proposed-*-excel.json"))
        got = json.loads(prop.read_text())["items"][0]
        self.assertEqual((got["op"], got["table"]), ("new", "decisions"))
        self.assertIn("invalid", got["flags"])
        self.assertIn("needs its date", got["invalid_reason"])

    def test_a_new_pending_decision_and_a_new_promise_from_excel(self):
        def decisions(ws, head):
            row = ws.max_row + 1
            for col, value in (("Project", "ABC"), ("Decision", "Choose the pilot venue"), ("Status", "Pending"),
                               ("Made by", "Sara Example")):
                ws.cell(row=row, column=head.index(col) + 1, value=value)
        self.edit("Decisions", decisions)

        def promises(ws, head):
            row = ws.max_row + 1
            for col, value in (("Project", "ABC"), ("Promised by", "Sara Example"), ("Promised to", "Alan Carter"),
                               ("Promise", "Send the data dictionary")):
                ws.cell(row=row, column=head.index(col) + 1, value=value)
        self.edit("Promises", promises)
        out = self.v.run("import", script="excel_sync.py").stdout
        self.assertIn("2 new", out)
        self.approve_and_apply()
        self.assertEqual(sql(self.v, "SELECT id, status, date FROM decisions WHERE decision = 'Choose the pilot venue'"),
                         [{"id": "ABC-D-006", "status": "Pending", "date": None}])
        self.assertEqual(sql(self.v, "SELECT id, promised_by FROM promises ORDER BY id")[-1],
                         {"id": "ABC-P-007", "promised_by": "Sara Example"})


# ── People labels (BL-94) ─────────────────────────────────────────────────────

class PeopleLabelTests(unittest.TestCase):
    def test_the_labels_column_is_read_and_written(self):
        sys.path.insert(0, str(REPO / "system/scripts"))
        import propose
        text = ("# People register\n\n## People\n\n| Full name | Email | Aliases | Organisation | Role | Hidden | Labels |\n"
                "|---|---|---|---|---|---|---|\n"
                "| Jane Doe | jane@client.example | JD | Client Ltd | Sponsor |  | Client principal, Executive sponsor |\n"
                "| Omar Ali |  |  | Enable | PM |  |  |\n")
        people = propose.people_parse(text)
        self.assertEqual([p["labels"] for p in people], [["Client principal", "Executive sponsor"], []])
        self.assertEqual(people[0]["extra"], {})                               # not an unknown extra column any more
        people[1]["labels"] = ["External"]
        again = propose.people_parse(propose.people_text(text, people))
        self.assertEqual([p["labels"] for p in again], [["Client principal", "Executive sponsor"], ["External"]])

    def test_a_register_without_the_column_gets_it_on_the_next_write(self):
        sys.path.insert(0, str(REPO / "system/scripts"))
        import propose
        old = "## People\n\n| Full name | Email |\n|---|---|\n| Jane Doe | j@x.example |\n"
        out = propose.people_text(old, propose.people_parse(old))
        self.assertIn("Labels", out.splitlines()[2])
        self.assertEqual(propose.people_parse(out)[0]["labels"], [])

    def test_the_template_and_sample_register_have_the_column(self):
        for f in ("system/templates/people-register-template.md", "projects/_portfolio/people-register.md"):
            header = next(ln for ln in (REPO / f).read_text(encoding="utf-8").splitlines() if ln.startswith("| Full name"))
            self.assertTrue(header.rstrip(" |").endswith("Labels"), f)             # after Hidden (and VIP, BL-99)


if __name__ == "__main__":
    unittest.main()
