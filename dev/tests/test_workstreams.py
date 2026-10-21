"""Tests for the workstreams sprint: the workstreams table (BL-58) and person pages, phase 2 (BL-56).

Run: python3 dev/tests/test_workstreams.py
Each class builds a throwaway vault with fictional data and runs the real scripts there.
"""
import json
import sqlite3
import subprocess
import sys
import unittest
from contextlib import closing
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from test_tracker import REPO, Vault, action  # noqa: E402


def ws(n, project="ABC", **v):
    return {"n": n, "op": "new", "table": "workstreams", "values": {"project": project, "source": "Alan (cx-edit)", **v},
            "match": "strong", "flags": [], "decision": "approved"}


def sql(v, query, *args):
    with closing(sqlite3.connect(v.root / "system/tracker.db")) as con:
        con.row_factory = sqlite3.Row
        return [dict(r) for r in con.execute(query, args)]


def add_project(v, code, name):
    """A second project in the register and the tracker (fictional)."""
    reg = v.root / "projects/_portfolio/project-register.md"
    reg.write_text(reg.read_text().replace("|  |  |  |  |  |  |  |  |",
                                           f"| {code} | {name} | {code} {name} | | | | | |"))
    (v.root / "projects" / f"{code} {name}").mkdir(parents=True, exist_ok=True)


class MigrationTests(unittest.TestCase):
    """Seeded from items and from settings; running it twice changes nothing; nothing is lost."""

    def setUp(self):
        self.v = v = Vault()
        add_project(v, "DEF", "Digital Enablement")
        # The schema before BL-58: today's, without the workstreams table. Read from the checked-out files, never a
        # git branch (a release build checks out a tag, with no local "main").
        import re
        old_schema = re.sub(r"-- Workstreams \(BL-58\).*?\n\);\n\n", "",
                            (REPO / "system/scripts/schema.sql").read_text(encoding="utf-8"), flags=re.S)
        self.assertNotIn("CREATE TABLE IF NOT EXISTS workstreams", old_schema)
        (v.root / "system/scripts/schema.sql").write_text(old_schema)
        v.run("init")
        (v.root / "system/scripts/schema.sql").write_text((REPO / "system/scripts/schema.sql").read_text())
        with closing(sqlite3.connect(v.root / "system/tracker.db")) as con:
            for i, (p, w) in enumerate([("ABC", "Data"), ("ABC", "data"), ("ABC", "Cut-over"), ("DEF", "Cut-over"),
                                        ("ABC", " "), ("ABC", None)], start=1):
                con.execute("INSERT INTO actions (id, project, description, workstream, raised, source) VALUES "
                            "(?, ?, ?, ?, '2026-10-01', 'test')", (f"{p}-A-{i:03d}", p, f"Item {i}", w))
            con.execute("INSERT INTO raid (id, project, type, description, workstream, raised, source) VALUES "
                        "('DEF-001', 'DEF', 'Risk', 'r', 'Testing', '2026-10-01', 'test')")
            con.commit()
        cfg = v.root / "system/config/settings.json"
        data = json.loads(cfg.read_text())
        data["workstreams"] = {"added": ["ABC|Commercial", "Legacy plain"], "hidden": ["DEF|Testing", "ABC|Retired one"]}
        cfg.write_text(json.dumps(data, indent=2))

    def tearDown(self):
        self.v.cleanup()

    def test_seeded_once_from_items_and_settings(self):
        before = sql(self.v, "SELECT * FROM actions ORDER BY id")
        out = self.v.run("migrate").stdout
        self.assertIn("Workstreams added from items and settings (6)", out)
        got = {(w["project"], w["name"]): (w["id"], w["status"], w["lead"]) for w in sql(self.v, "SELECT * FROM workstreams")}
        self.assertEqual(got, {("ABC", "Commercial"): ("ABC-W-001", "Active", None),
                               ("ABC", "Cut-over"): ("ABC-W-002", "Active", None),
                               ("ABC", "Data"): ("ABC-W-003", "Active", None),              # "data" is the same one
                               ("ABC", "Retired one"): ("ABC-W-004", "Closed", None),
                               ("DEF", "Cut-over"): ("DEF-W-001", "Active", None),          # same name, its own row
                               ("DEF", "Testing"): ("DEF-W-002", "Closed", None)})          # hidden in settings
        self.assertEqual(sql(self.v, "SELECT * FROM actions ORDER BY id"), before)          # nothing lost or changed
        snapshot = sql(self.v, "SELECT * FROM workstreams ORDER BY id")
        out2 = self.v.run("migrate").stdout
        self.assertNotIn("Workstreams added", out2)
        self.assertEqual(sql(self.v, "SELECT * FROM workstreams ORDER BY id"), snapshot)      # twice: no change


class TrackerTests(unittest.TestCase):
    """Validation, and new, rename and close as proposals, with undo."""

    def setUp(self):
        self.v = v = Vault()
        add_project(v, "DEF", "Digital Enablement")
        v.run("init")
        v.propose("ws", [ws(1, name="Data", lead="Sara Example", keywords="extract, masking", aliases="DQ"),
                         ws(2, name="Testing", aliases="UAT"), ws(3, project="DEF", name="Data")])
        v.run("apply")

    def tearDown(self):
        self.v.cleanup()

    def test_items_need_an_active_workstream_of_their_project(self):
        v = self.v
        bad = v.propose("bad", [action(1, workstream="Commercial"), action(2, project="DEF", workstream="Testing")])
        out = v.run("validate", str(bad), ok=False).stdout
        self.assertIn("“Commercial” isn't an active workstream of ABC", out)
        self.assertIn("“Testing” isn't an active workstream of DEF", out)              # ABC's Testing doesn't count for DEF
        self.assertIn("Nothing was written", v.run("apply", ok=False).stdout)          # one bad item stops the apply
        bad.unlink()
        v.propose("good", [action(1, description="Any case", workstream="data"), action(2, description="Blank", workstream="")])
        v.run("apply")
        self.assertEqual({a["description"]: a["workstream"] for a in v.rows("actions")}, {"Any case": "Data", "Blank": ""})   # stored as the workstream spells it

    def test_duplicate_name_refused_any_case(self):
        self.v.propose("dup", [ws(1, name="testing ")])
        self.assertIn("already has a workstream called “testing”", self.v.run("apply", ok=False).stdout)

    def test_rename_carries_items_and_undo_reverses_it(self):
        v = self.v
        v.propose("items", [action(1, description="T", workstream="Testing"), action(2, description="D", workstream="Data"),
                            action(3, project="DEF", description="D", workstream="Data")])
        v.run("apply")
        before = (v.rows("actions"), v.rows("workstreams"))
        abc_data = next(w for w in v.rows("workstreams") if w["project"] == "ABC" and w["name"] == "Data")
        v.propose("rename", [{"n": 1, "op": "update", "table": "workstreams", "id": abc_data["id"],
                              "changes": {"name": "Data and analytics"}, "decision": "approved"}])
        v.run("apply")
        acts = {a["project"] + " " + a["description"]: a["workstream"] for a in v.rows("actions")}
        self.assertEqual(acts, {"ABC T": "Testing", "ABC D": "Data and analytics", "DEF D": "Data"})   # DEF's Data untouched
        self.assertEqual(sorted(w["name"] for w in v.rows("workstreams")), ["Data", "Data and analytics", "Testing"])
        v.run("undo")
        self.assertEqual((v.rows("actions"), v.rows("workstreams")), before)

    def test_close_blocks_new_use_only_and_undo_reopens(self):
        v = self.v
        v.propose("items", [action(1, workstream="Testing")])
        v.run("apply")
        testing = next(w for w in v.rows("workstreams") if w["name"] == "Testing")
        v.propose("close", [{"n": 1, "op": "cancel", "table": "workstreams", "id": testing["id"], "decision": "approved"}])
        v.run("apply")
        self.assertEqual(next(w for w in v.rows("workstreams") if w["name"] == "Testing")["status"], "Closed")
        a1 = v.rows("actions")[0]
        v.propose("edit", [{"n": 1, "op": "update", "table": "actions", "id": a1["id"], "changes": {"priority": "High"},
                            "decision": "approved"}])
        v.run("apply")                                                                     # an item already on it is fine
        v.propose("more", [action(2, workstream="Testing")])
        self.assertIn("isn't an active workstream", v.run("apply", ok=False).stdout)        # new use is refused
        v.run("undo")                                                                      # the priority edit
        v.run("undo")                                                                      # the close
        self.assertEqual(next(w for w in v.rows("workstreams") if w["name"] == "Testing")["status"], "Active")

    def test_a_workstream_stays_in_its_project(self):
        w = next(w for w in self.v.rows("workstreams") if w["project"] == "ABC" and w["name"] == "Data")
        self.v.propose("move", [{"n": 1, "op": "update", "table": "workstreams", "id": w["id"],
                                 "changes": {"project": "DEF"}, "decision": "approved"}])
        self.assertIn("a workstream stays in its project", self.v.run("apply", ok=False).stdout)


class ExtractionAndReviewTests(unittest.TestCase):
    """Matching only within the project; an unknown workstream is left blank and flagged; Review adds or skips it."""

    def setUp(self):
        self.v = v = Vault()
        add_project(v, "DEF", "Digital Enablement")
        v.run("init")
        v.propose("ws", [ws(1, name="Data", aliases="DQ, data quality"), ws(2, project="DEF", name="Integration")])
        v.run("apply")

    def tearDown(self):
        self.v.cleanup()

    def staged(self):
        p = max((self.v.root / "system/staging").glob("proposed-*.json"), key=lambda f: f.stat().st_mtime)
        return p, json.loads(p.read_text())

    def test_minutes_and_sweep_match_within_the_project(self):
        v = self.v
        spec = v.root / "system/staging/drafts/m.json"
        spec.write_text(json.dumps({"date": "2026-10-06", "title": "ABC Steering 6", "projects": ["ABC"], "items": [
            {"table": "actions", "values": {"project": "ABC", "description": "Profile the extract", "workstream": "dq"}},
            {"table": "actions", "values": {"project": "ABC", "description": "Book the API test", "workstream": "Integration"}},
            {"table": "raid", "values": {"project": "ABC", "type": "Risk", "description": "Late data", "workstream": "Data"}}]}))
        v.run("stage", str(spec), script="minutes.py")
        _, data = self.staged()
        got = [(i["values"].get("workstream"), i.get("ws_suggest"), "new-workstream" in i["flags"]) for i in data["items"]]
        self.assertEqual(got, [("Data", None, False),                 # an alias, stored as the workstream's name
                               ("", "Integration", True),             # DEF's Integration doesn't count for ABC
                               ("Data", None, False)])
        sweep = v.root / "system/staging/drafts/s.json"
        sweep.write_text(json.dumps({"window": {}, "items": [{"table": "actions", "values": {
            "project": "DEF", "description": "Agree the interface", "workstream": "integration",
            "source": "Email · From: Sam · Link: https://outlook/x"}}]}))
        v.run("stage", str(sweep), script="sweep.py")
        _, data = self.staged()
        self.assertEqual(data["items"][0]["values"]["workstream"], "Integration")

    def test_review_adds_or_skips_the_suggestion(self):
        v = self.v
        v.propose("from-minutes", [{**action(1, description="Book the API test", workstream=""), "decision": "pending",
                                    "flags": ["new-workstream"], "ws_suggest": "Integration"},
                                   {**action(2, description="Second", workstream=""), "decision": "pending",
                                    "flags": ["new-workstream"], "ws_suggest": "Reporting"}])
        listed = v.run("list", script="review.py").stdout
        self.assertIn("New workstream “Integration” for ABC?", listed)
        nums = {r["text"]: r["n"] for r in json.loads(v.run("list", "--json", script="review.py").stdout)}
        out = v.run("workstream", str(nums["Book the API test"]), "add", "--lead=Sara Example", "--keywords=API",
                    script="review.py").stdout
        self.assertIn("added to ABC and set", out)
        v.run("workstream", str(nums["Second"]), "skip", script="review.py")
        data = json.loads((v.root / "system/staging/proposed-from-minutes.json").read_text())
        self.assertEqual([(i["table"], i["decision"]) for i in data["items"]],
                         [("workstreams", "approved"), ("actions", "pending"), ("actions", "pending")])   # before the item
        self.assertEqual(data["items"][1]["values"]["workstream"], "Integration")
        self.assertEqual((data["items"][2]["values"]["workstream"], data["items"][2]["flags"]), ("", []))
        v.run("list", script="review.py")
        nums = {r["text"]: r["n"] for r in json.loads(v.run("list", "--json", script="review.py").stdout)}
        v.run("approve", f"{nums['Book the API test']},{nums['Second']}", script="review.py")
        v.run("apply")
        w = next(w for w in v.rows("workstreams") if w["name"] == "Integration" and w["project"] == "ABC")
        self.assertEqual((w["lead"], w["keywords"], w["status"]), ("Sara Example", "API", "Active"))
        self.assertEqual(sorted((a["description"], a["workstream"]) for a in v.rows("actions")),
                         [("Book the API test", "Integration"), ("Second", "")])


class ExcelAndEditTests(unittest.TestCase):
    """The Workstreams sheet round-trips; cx-edit suggests, adds and retires the old settings lists."""

    def setUp(self):
        self.v = v = Vault()
        v.run("init")
        v.propose("ws", [ws(1, name="Testing")])
        v.run("apply")

    def tearDown(self):
        self.v.cleanup()

    def test_excel_round_trip(self):
        sys.path.insert(0, str(REPO / "system/scripts/vendor"))
        from openpyxl import load_workbook
        v = self.v
        v.run("export", script="excel_sync.py")
        book = v.root / "tracker-view.xlsx"
        wb = load_workbook(book)
        sheet = wb["Workstreams"]
        head = [c.value for c in sheet[1]]
        self.assertEqual(head[:4], ["ID", "Project", "Workstream", "Lead"])
        sheet.cell(row=2, column=head.index("Aliases") + 1, value="UAT")
        sheet.append(["", "ABC", "Data"])
        wb.save(book)
        v.run("import", script="excel_sync.py")
        p = next((v.root / "system/staging").glob("proposed-*excel*.json"))
        data = json.loads(p.read_text())
        for i in data["items"]:
            i["decision"] = "approved"
        p.write_text(json.dumps(data))
        v.run("apply")
        self.assertEqual({w["name"]: w["aliases"] for w in v.rows("workstreams")}, {"Testing": "UAT", "Data": None})

    def test_suggest_add_and_retire(self):
        v = self.v
        v.propose("items", [action(1, description="Run UAT regression pack", owner="Sara Example", workstream="Testing",
                                   source="Teams · Channel: ABC – Testing · Sent: 2026-10-02"),
                            action(2, description="Fix UAT regression defects", owner="Sara Example", workstream="Testing")])
        v.run("apply")
        s = json.loads(v.run("workstreams", "suggest", "ABC", "Testing", script="propose.py").stdout)
        self.assertEqual((s["items"], s["aliases"], s["channel"]), (2, "UAT", "ABC – Testing"))
        self.assertIn("regression", s["keywords"])
        self.assertEqual(s["lead"], "Sara Example")
        v.run("workstreams", "add", "ABC", "Reporting", "--keywords", "dashboards", script="propose.py")
        v.run("apply")
        self.assertIn("Reporting", [w["name"] for w in v.rows("workstreams")])
        cfg = v.root / "system/config/settings.json"
        d = json.loads(cfg.read_text())
        d["workstreams"] = {"added": ["ABC|Reporting", "ABC|Gone"], "hidden": []}
        cfg.write_text(json.dumps(d))
        self.assertIn("Not yet in the workstreams table: ABC|Gone",
                      v.run("workstreams", "retire-settings", script="propose.py", ok=False).stdout)
        v.run("migrate")
        v.run("workstreams", "retire-settings", script="propose.py")
        self.assertEqual(json.loads(cfg.read_text())["workstreams"], {"added": [], "hidden": []})


FAKE_VAULT = """const FILES={}; VW={read:async p=>{ if(!(p in FILES)) throw new Error("not found: "+p); return FILES[p]; },
  write:async(p,t)=>{ FILES[p]=t; }, remove:async p=>{ delete FILES[p]; }, exists:async p=>p in FILES,
  append:async(p,t)=>{ FILES[p]=(FILES[p]||"")+t; }, upload:async()=>{} }; VSTATE="ok"; needVault=async()=>true;
  const wait=ms=>new Promise(r=>setTimeout(r,ms)), set=(id,v)=>{ const e=document.getElementById(id); e.value=v; e.dispatchEvent(new Event("input",{bubbles:true})); };"""


class DashboardTests(unittest.TestCase):
    """The picker, the project's settings › Workstreams (Settings › Workstreams until the Settings line), the project page
    card and Review in headless Chrome; the proposals the
    page writes are then applied by the real tracker."""

    @classmethod
    def setUpClass(cls):
        from test_dashboard import CHROME, populate
        if not CHROME.exists():
            raise unittest.SkipTest("Google Chrome not installed")
        from test_insight import chrome
        cls.v = v = Vault()
        populate(v)
        v.propose("ws", [ws(1, name="Data", lead="Sara Example"), ws(2, name="Testing", aliases="UAT")])
        v.run("apply")
        v.propose("items", [action(1, description="Profile the data extract", owner="Sara Example", workstream="Data",
                                   source="Teams · Channel: ABC – Data · Sent: 2026-10-02")])
        v.run("apply")
        v.propose("flagged", [{**action(1, description="Book the API test", workstream=""), "decision": "pending",
                               "flags": ["new-workstream"], "ws_suggest": "Integration"}])
        v.run(script="build_dashboard.py")
        cls.out, cls.errors = chrome(v, FAKE_VAULT + """
          go("cockpit","ABC"); await wait(30); const card=document.getElementById("wsCard")?document.getElementById("wsCard").innerText:null, cockpit=document.getElementById("v-cockpit").innerText;
          // the picker (BL-179: a new item is the peek, as an existing one; BL-187: a name the project doesn't have is added from
          // the list, with its details, in the item's own proposal)
          openNew("actions"); const nid=CUR.nid; await pkCommit("actions",nid,{p:"ABC"},"Project","",""); await pkCommit("actions",nid,{t:"Draft the reporting pack"},"Action","","");
          PKPROP.ws(document.getElementById("pkTitle"),CUR.draft,"actions",nid); const wq=document.querySelector(".pkpop .pkq"); wq.value="Reporting"; wq.dispatchEvent(new Event("input"));
          [...document.querySelectorAll(".pkpop .pkopt")].find(o=>o.textContent.startsWith("Add “Reporting”")).click(); await wait(20);
          const form=[document.querySelector(".pkws .pkws-h").textContent,document.querySelector('.pkws [data-wsn="name"]').value];
          document.querySelector('.pkws [data-wsn="keywords"]').value="pack"; document.querySelector('.pkws [data-wsf="add"]').click(); await wait(30); await pkQuiet(); await VQ;
          const ask=document.getElementById("pkNote").textContent, box=CUR.draft.ws, offered=wsFor("ABC"); closePanel();
          // the project's settings › Workstreams (#247; the old Settings › Workstreams address opens it from a project): add, rename, retire
          go("cockpit","ABC"); go("settings","workstreams"); await wait(60); const wsAt=[VIEW,psCode.value,psCat.value];
          const row=()=>document.querySelector('#psWs tr.wsedit'), fill=(f,v)=>{ row().querySelector(`[data-wsf="${f}"]`).value=v; };
          document.getElementById("psAddWs").click(); await wait(30); fill("name","Integration"); fill("keywords","api, interfaces");
          row().querySelector("[data-ws-save]").click(); await wait(80); await VQ; await wait(30);
          document.querySelector('#psWs [data-ws-edit="Data"]').click(); await wait(30); fill("name","Data and analytics");
          row().querySelector("[data-ws-save]").click(); await wait(80); await VQ; await wait(30);
          document.querySelector('#psWs [data-ws-edit="Testing"]').click(); await wait(30); row().querySelector("[data-ws-retire]").click(); await wait(80); await VQ; await wait(30);
          const mgr=document.getElementById("psWs").innerText, renamedItem=A.find(a=>a.t==="Profile the data extract").ws;
          // Intake: the flagged item offers to add its workstream
          go("intake-review"); await wait(150); const q=document.querySelector("#v-intake-review .wsq")?.innerText||"";
          window.OUT={card, cockpit, form, ask, box, offered, wsAt, mgr, renamedItem, q,
            files:Object.fromEntries(Object.entries(FILES).filter(([k])=>k.startsWith("system/staging/proposed-")))};""")

    @classmethod
    def tearDownClass(cls):
        cls.v.cleanup()

    def test_page(self):
        o = self.out
        self.assertEqual(self.errors, [])
        self.assertIsNone(o["card"])                                                      # BL-69: no Workstreams card on the project page
        self.assertNotIn("Manage workstreams", o["cockpit"])
        self.assertNotIn("Workstreams", o["cockpit"])
        self.assertEqual(o["form"], ["New workstream for Example Programme", "Reporting"])     # BL-187: added from the peek
        self.assertIn("New workstream “Reporting” for Example Programme", o["ask"])
        self.assertEqual(o["box"], "Reporting")
        self.assertIn("Reporting", o["offered"])
        self.assertEqual(o["wsAt"], ["psettings", "ABC", "ws"])                         # the old address opens the project's Workstreams
        self.assertIn("Data and analytics", o["mgr"])
        self.assertNotIn("Testing", o["mgr"])                                              # retired: no longer listed
        self.assertEqual(o["renamedItem"], "Data and analytics")                          # the item follows the rename
        self.assertIn("New workstream “Integration” for ABC?", o["q"])

    def test_proposals_apply(self):
        v = self.v
        files = self.out["files"]
        first = next(d for d in map(json.loads, files.values()) if any(i["table"] == "actions" for i in d["items"]))
        self.assertEqual([(i["table"], i["op"]) for i in first["items"]], [("workstreams", "new"), ("actions", "new")])   # BL-187: ahead of it
        self.assertEqual(first["items"][0]["values"]["keywords"], "pack")
        for k, t in files.items():
            (v.root / k).write_text(t)
        (v.root / "system/staging/proposed-flagged.json").unlink()
        v.run("apply")
        got = {w["name"]: (w["status"], w["channel"]) for w in v.rows("workstreams")}
        self.assertEqual(got, {"Data and analytics": ("Active", None), "Testing": ("Closed", None),
                               "Integration": ("Active", None), "Reporting": ("Active", None)})
        # the project settings page has no Teams channel field (#247): a channel is set in the Excel sheet or with cx-edit
        self.assertEqual({w["name"]: w["keywords"] for w in v.rows("workstreams")}["Integration"], "api, interfaces")
        acts = {a["description"]: a["workstream"] for a in v.rows("actions")}
        self.assertEqual((acts["Draft the reporting pack"], acts["Profile the data extract"]), ("Reporting", "Data and analytics"))


def key_person_vault():
    """ABC with 7 open items, 5 of them Sara's; DEF with 3 (under the minimum)."""
    v = Vault()
    add_project(v, "DEF", "Digital Enablement")
    v.run("init")
    items = [action(i, description=f"Sara task {i}", owner="Sara Example") for i in range(1, 6)] + \
            [action(6, description="Omar task", owner="Omar Ali"), action(7, description="Nobody yet", owner="TBC")] + \
            [action(8 + i, project="DEF", description=f"DEF {i}", owner="Omar Ali") for i in range(3)]
    v.propose("items", items)
    v.run("apply")
    return v


class KeyPersonTests(unittest.TestCase):
    """BL-56: one person owning most of a project's open items; thresholds, switch, health evidence, never client."""

    def setUp(self):
        self.v = key_person_vault()

    def tearDown(self):
        self.v.cleanup()

    def warnings(self):
        out = subprocess.run([sys.executable, "-c", "import sys, json; sys.path.insert(0, 'system/scripts')\n"
                              "import paths, build_dashboard as b\nfrom datetime import date\n"
                              "print(json.dumps([w for w in b.early_warnings(paths.load_settings(), date.today()) "
                              "if w['sig'] == 'key_person']))"], cwd=self.v.root, capture_output=True, text=True)
        return json.loads(out.stdout)

    def set_kp(self, **v):
        cfg = self.v.root / "system/config/settings.json"
        d = json.loads(cfg.read_text())
        d["warnings"]["key_person"].update(v)
        cfg.write_text(json.dumps(d))

    def test_fires_and_thresholds(self):
        self.assertEqual(self.warnings(), [{"k": "projects", "id": "ABC", "p": "ABC", "sig": "key_person",
                                            "who": "Sara Example", "text": "Sara Example owns 5 of 7 open items (71%)"}])
        self.set_kp(share=75)
        self.assertEqual(self.warnings(), [])
        self.set_kp(share=50, min_items=8)
        self.assertEqual(self.warnings(), [])
        self.set_kp(min_items=6, on=False)
        self.assertEqual(self.warnings(), [])

    def test_in_health_evidence_never_in_client_report(self):
        ev = subprocess.run([sys.executable, "-c", "import sys, json; sys.path.insert(0, 'system/scripts')\n"
                             "import report\nprint(json.dumps(report.health_evidence('ABC')))"], cwd=self.v.root,
                            capture_output=True, text=True).stdout
        self.assertIn("Example Programme: Sara Example owns 5 of 7 open items (71%)", json.loads(ev)["warnings"])
        client = self.v.run("client-data", "ABC", "--since", "2026-09-01", script="report.py").stdout
        self.assertNotIn("owns 5 of 7", client)


class PersonPageTests(unittest.TestCase):
    """The Draft email button (BL-83, in place of Prep 1:1), handover (one proposal, undone in one go), flags and items closed lately, in headless Chrome."""

    @classmethod
    def setUpClass(cls):
        from test_dashboard import CHROME
        if not CHROME.exists():
            raise unittest.SkipTest("Google Chrome not installed")
        from test_insight import chrome
        from datetime import date, timedelta
        cls.v = v = key_person_vault()
        done = (date.today() - timedelta(days=3)).isoformat()
        v.propose("closed", [{**action(1, description="Sara closed one", owner="Sara Example", status="Closed",
                                       closed=done, raised="2026-09-01"), "n": 1}])
        v.propose("ws", [ws(2, name="Data", lead="Sara Example")])
        v.run("apply")
        v.run(script="build_dashboard.py")
        cls.out, cls.errors = chrome(v, FAKE_VAULT + """
          go("person","Sara Example"); await wait(30); const page=document.getElementById("v-person").innerText;
          const buttons=[...document.querySelectorAll("#v-person .hhead button, #v-person button.btn")].map(b=>b.textContent.trim()).filter(Boolean);
          const gone={pp1:!!document.getElementById("pp1"), p1Dlg:!!document.getElementById("p1Dlg"), openP1:typeof openP1, p1Dlg_fn:typeof p1Dlg};
          document.getElementById("ppHand").click(); await wait(20);
          const lis=[...document.querySelectorAll("#v-person .handpick")]; lis[0].click(); await wait(10);
          document.querySelectorAll("#v-person .handpick")[1].click(); await wait(10);
          document.getElementById("handTo").value="Omar Ali"; document.getElementById("handGo").click(); await wait(150); await VQ;
          const after=A.filter(a=>a.o==="Omar Ali"&&a.p==="ABC").length;
          window.OUT={page, buttons, gone, after, files:Object.fromEntries(Object.entries(FILES).filter(([k])=>k.startsWith("system/staging/proposed-")))};""")

    @classmethod
    def tearDownClass(cls):
        cls.v.cleanup()

    def test_page_and_the_draft_email_button(self):
        o = self.out
        self.assertEqual(self.errors, [])
        self.assertNotIn("Prep 1:1", o["page"])                                              # replaced by Draft email (BL-83)
        self.assertIn("Draft email", o["buttons"])
        self.assertLess(o["buttons"].index("Draft email"), o["buttons"].index("Hand over…"))   # beside Hand over, before it
        self.assertEqual(o["gone"], {"pp1": False, "p1Dlg": False, "openP1": "undefined", "p1Dlg_fn": "undefined"})
        self.assertIn("Example Programme: Sara Example owns 5 of 7 open items (71%)", o["page"])       # the key-person note
        self.assertIn("Closed in the last 14 days", o["page"])
        self.assertIn("Sara closed one", o["page"])

    def test_handover_is_one_proposal_and_undoes_in_one_go(self):
        v = self.v
        self.assertEqual(self.out["after"], 3)                                             # Omar had 1 in ABC, now 3
        files = self.out["files"]
        self.assertEqual(len(files), 1)
        data = json.loads(next(iter(files.values())))
        self.assertEqual([(i["table"], i["op"], i["changes"]) for i in data["items"]],
                         [("actions", "update", {"owner": "Omar Ali"})] * 2)
        before = v.rows("actions")
        for k, t in files.items():
            (v.root / k).write_text(t)
        v.run("apply")
        self.assertEqual(sum(a["owner"] == "Omar Ali" and a["project"] == "ABC" for a in v.rows("actions")), 3)
        v.run("undo")
        self.assertEqual(v.rows("actions"), before)


if __name__ == "__main__":
    unittest.main(verbosity=2)
