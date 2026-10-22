"""BL-186 stage A (sprint-milestones): the milestone data under change control. The five states with Alan's state and its reason,
owner, the achieved date, client wording, the rebaseline's decision, the Missed migration, propose, Excel, undo and recode.

Run: python3 dev/tests/test_milestone_states.py
"""
import json
import sqlite3
import subprocess
import sys
import unittest
from contextlib import closing
from datetime import date, timedelta
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from test_tracker import Vault  # noqa: E402
from test_dashboard import populate  # noqa: E402

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "system" / "scripts" / "vendor"))
from openpyxl import load_workbook  # noqa: E402

TODAY = date.today()


def iso(n):
    return (TODAY + timedelta(days=n)).isoformat()


def milestone(n, name, base, fc=None, **v):
    return {"n": n, "op": "new", "table": "milestones", "match": "strong", "flags": [], "decision": "approved",
            "values": {"project": "ABC", "name": name, "baseline_date": base, "forecast_date": fc or base,
                       "source": "Alan (cx-edit), 2026-10-05", **v}}


def upd(id_, n=1, rebaseline=None, **changes):
    it = {"n": n, "op": "update", "table": "milestones", "id": id_, "changes": changes, "match": "strong", "flags": [],
          "decision": "approved"}
    if rebaseline:
        it["rebaseline"] = rebaseline
    return it


# The milestones table as 1.4.0 left it (Missed allowed, none of the BL-186 columns)
OLD_MILESTONES = """CREATE TABLE milestones (
  id            TEXT PRIMARY KEY CHECK (id GLOB '[A-Z]*-M-[0-9][0-9][0-9]' AND length(id) BETWEEN 9 AND 13 AND substr(id, 1, length(id) - 6) NOT GLOB '*[^A-Z0-9]*'),
  project       TEXT NOT NULL REFERENCES projects(code),
  name          TEXT NOT NULL,
  description   TEXT,
  baseline_date TEXT CHECK (baseline_date IS NULL OR baseline_date GLOB '[0-9][0-9][0-9][0-9]-[0-9][0-9]-[0-9][0-9]'),
  forecast_date TEXT CHECK (forecast_date IS NULL OR forecast_date GLOB '[0-9][0-9][0-9][0-9]-[0-9][0-9]-[0-9][0-9]'),
  status        TEXT NOT NULL DEFAULT 'Planned' CHECK (status IN ('Planned','Done','Missed','Cancelled')),
  sort_order    INTEGER NOT NULL DEFAULT 0,
  percent_complete INTEGER CHECK (percent_complete IS NULL OR percent_complete BETWEEN 0 AND 100),
  source        TEXT NOT NULL,
  updated_at    TEXT NOT NULL DEFAULT (datetime('now'))
)"""
OLD_COLS = "id, project, name, description, baseline_date, forecast_date, status, sort_order, percent_complete, source, updated_at"


class Case(unittest.TestCase):
    def setUp(self):
        self.v = Vault()
        self.v.run("init")
        self.v.propose("m", [milestone(1, "UAT complete", iso(30)), milestone(2, "Go-live", iso(60))])
        self.v.run("apply")

    def tearDown(self):
        self.v.cleanup()

    def ms(self, id_="ABC-M-001"):
        return next(m for m in self.v.rows("milestones") if m["id"] == id_)

    def apply(self, name, items, ok=True):
        self.v.propose(name, items)
        return self.v.run("apply", ok=ok)

    def invalid(self, name, items):
        p = self.v.propose(name, items)
        out = self.v.run("validate", str(p), ok=False).stdout
        p.unlink()                                       # a refused proposal would hold back the next apply
        return out


class Fields(Case):
    """Each new field through propose (preview and stage), apply and undo."""

    def test_each_field_applies_and_undoes(self):
        cases = [
            ({"owner": "Sara Example"}, {"owner": "Sara Example"}),
            ({"state": "At risk"}, {"state": "At risk", "state_set": 1, "state_reason": None}),
            ({"state": "At risk", "state_reason": "Security review open 23 days"}, {"state": "At risk", "state_reason": "Security review open 23 days"}),
            ({"client_visible": 1, "client_message": "Testing starts once data access is granted."},
             {"client_visible": 1, "client_message": "Testing starts once data access is granted."}),
            ({"status": "Done", "actual_date": iso(-2), "percent_complete": 100},
             {"status": "Done", "actual_date": iso(-2), "percent_complete": 100, "state": None}),
            ({"percent_complete": 40}, {"percent_complete": 40}),
            ({"internal_note": "Sponsor call weekly until go-live"}, {"internal_note": "Sponsor call weekly until go-live"}),
        ]
        for i, (changes, want) in enumerate(cases):
            before = self.ms()
            out = self.apply(f"f{i}", [upd("ABC-M-001", **changes)]).stdout
            self.assertIn("written 1", out)
            got = self.ms()
            for k, val in want.items():
                self.assertEqual(got[k], val, (changes, k))
            self.v.run("undo")
            after = self.ms()
            self.assertEqual({k: after[k] for k in before if k != "updated_at"}, {k: before[k] for k in before if k != "updated_at"}, changes)

    def test_propose_previews_and_stages_a_state_with_the_dates_view(self):
        spec = self.v.root / "spec.json"
        spec.write_text(json.dumps({"items": [{"op": "update", "table": "milestones", "id": "ABC-M-001",
                                               "changes": {"state": "Slipped", "owner": "Sara Example"}}]}))
        out = self.v.run("preview", str(spec), script="propose.py").stdout
        self.assertIn("State: — → Slipped", out)
        self.assertIn("Owner: — → Sara Example", out)
        self.assertIn("The dates suggest On track", out)
        self.v.run("stage", str(spec), script="propose.py")
        self.v.run("apply")
        self.assertEqual((self.ms()["state"], self.ms()["state_set"]), ("Slipped", 1))

    def test_clearing_the_state_returns_it_to_the_dates(self):
        self.apply("a", [upd("ABC-M-001", state="At risk")])
        self.apply("b", [upd("ABC-M-001", state="")])
        m = self.ms()
        self.assertEqual((m["state"], m["state_set"], m["state_reason"]), (None, 0, None))


class Rules(Case):
    def test_a_better_state_than_the_dates_needs_a_reason(self):
        self.apply("slip", [upd("ABC-M-001", forecast_date=iso(37))])          # the dates now say Slipped
        out = self.invalid("ok", [upd("ABC-M-001", state="On track")])
        self.assertIn("On track is better than the dates suggest (Slipped)", out)
        self.assertIn("state_reason", out)
        self.apply("why", [upd("ABC-M-001", state="On track", state_reason="Vendor confirmed recovery plan")])
        self.assertEqual(self.ms()["state"], "On track")
        self.apply("worse", [upd("ABC-M-002", state="At risk")])                # worse than the dates: no reason needed
        self.assertEqual(self.ms("ABC-M-002")["state"], "At risk")
        self.apply("past", [upd("ABC-M-002", forecast_date=iso(70))])
        out = self.invalid("past2", [upd("ABC-M-002", state="On track")])
        self.assertIn("better than the dates suggest (Slipped)", out)

    def test_overdue_is_suggested_once_the_agreed_date_has_passed(self):
        self.apply("old", [milestone(1, "Data access", iso(-3))])
        m = self.ms("ABC-M-003")
        sys.path.insert(0, str(self.v.root / "system/scripts"))
        import tracker
        self.assertEqual(tracker.suggested_state(m), "Overdue")
        self.assertEqual(tracker.milestone_state(m), "Overdue")
        out = self.invalid("slipped", [upd("ABC-M-003", state="Slipped")])
        self.assertIn("better than the dates suggest (Overdue)", out)

    def test_client_visible_needs_its_wording(self):
        out = self.invalid("cv", [upd("ABC-M-001", client_visible=1)])
        self.assertIn("needs its client wording", out)
        out = self.invalid("cv2", [upd("ABC-M-001", client_visible=1, client_message="  ")])
        self.assertIn("needs its client wording", out)
        self.apply("cv3", [upd("ABC-M-001", client_visible=1, client_message="On plan for the agreed date.")])
        out = self.invalid("cv4", [upd("ABC-M-001", client_message="")])
        self.assertIn("needs its client wording", out)

    def test_achieved_is_done_with_its_date(self):
        self.assertIn("an Achieved milestone is Done", self.invalid("a", [upd("ABC-M-001", state="Achieved")]))
        self.assertIn("only when the milestone is Done", self.invalid("b", [upd("ABC-M-001", actual_date=iso(-1))]))
        self.assertIn("can't be after today", self.invalid("c", [upd("ABC-M-001", status="Done", actual_date=iso(2))]))
        self.apply("d", [upd("ABC-M-001", status="Done")])
        self.assertEqual(self.ms()["actual_date"], TODAY.isoformat())             # Done fills today when no date is given
        self.apply("e", [upd("ABC-M-001", status="Planned")])
        self.assertIsNone(self.ms()["actual_date"])                                # reopening clears it
        self.apply("f", [upd("ABC-M-001", status="Done", state="Achieved", actual_date=iso(-4))])
        self.assertEqual((self.ms()["state"], self.ms()["actual_date"]), ("Achieved", iso(-4)))
        self.assertIn("a Done milestone is Achieved", self.invalid("g", [upd("ABC-M-001", state="At risk")]))

    def test_missed_is_refused(self):
        self.assertIn("Planned, Done, Cancelled", self.invalid("m", [upd("ABC-M-001", status="Missed")]))
        self.assertIn("state must be one of", self.invalid("s", [upd("ABC-M-001", state="Late")]))


class Rebaseline(Case):
    RB = {"agreed_with": ["Sara Example", "Omar Ali"], "reason": "Data access granted late", "source": "Steering minutes 2026-10-01"}

    def test_a_rebaseline_needs_who_why_and_where(self):
        out = self.invalid("rb", [upd("ABC-M-002", baseline_date=iso(80))])
        self.assertIn("is a rebaseline and records a decision", out)
        self.assertIn("agreed_with", out)
        out = self.invalid("rb2", [upd("ABC-M-002", baseline_date=iso(80), rebaseline={"agreed_with": "Sara Example"})])
        self.assertIn("the reason, the source", out)

    def test_a_rebaseline_makes_a_decision_and_undo_takes_both(self):
        self.apply("rb", [upd("ABC-M-002", baseline_date=iso(80), forecast_date=iso(80), rebaseline=self.RB)])
        m = self.ms("ABC-M-002")
        self.assertEqual(m["baseline_date"], iso(80))
        d = self.v.rows("decisions")
        self.assertEqual(len(d), 1)
        self.assertEqual((d[0]["made_by"], d[0]["reasoning"], d[0]["source"], d[0]["status"], d[0]["date"]),
                         ("Sara Example, Omar Ali", "Data access granted late", "Steering minutes 2026-10-01", "Made", TODAY.isoformat()))
        self.assertIn("agreed date moves from", d[0]["decision"])
        applied = json.loads(next((self.v.root / "system/staging/applied").glob("proposed-rb.json")).read_text())
        self.assertEqual(applied["items"][0]["rebaseline"]["decision"], d[0]["id"])
        notes = [r for r in self.v.rows("change_log") if r["field"] == "note" and r["item_id"] == "ABC-M-002"]
        self.assertTrue(notes and d[0]["id"] in notes[-1]["new_value"])
        sys.path.insert(0, str(self.v.root / "system/scripts"))
        import tracker
        with closing(tracker.connect(self.v.root / "system/tracker.db")) as con:
            self.assertEqual(tracker.milestone_original(con, "ABC-M-002"), iso(60))      # the Original is the first agreed date
        self.v.run("undo")
        self.assertEqual(self.v.rows("decisions"), [])
        self.assertEqual(self.ms("ABC-M-002")["baseline_date"], iso(60))

    def test_a_rebaseline_is_never_approved_as_routine(self):
        sys.path.insert(0, str(self.v.root / "system/scripts"))
        import review
        self.assertEqual(review.individual_only(upd("ABC-M-002", baseline_date=iso(80), rebaseline=self.RB)), "rebaseline")
        self.assertIsNone(review.individual_only(upd("ABC-M-002", forecast_date=iso(80))))

    def test_the_first_agreed_date_is_not_a_rebaseline(self):
        self.apply("nb", [milestone(1, "Hypercare", None, iso(90))])
        self.apply("set", [upd("ABC-M-003", baseline_date=iso(90))])
        self.assertEqual(self.v.rows("decisions"), [])


class Recode(Case):
    def test_a_recode_keeps_the_new_fields_and_the_history(self):
        import shutil
        from test_tracker import ABC
        self.apply("x", [upd("ABC-M-002", owner="Sara Example", state="At risk", client_visible=1,
                             client_message="Go-live depends on data access.")])
        self.apply("rb", [upd("ABC-M-002", baseline_date=iso(80), rebaseline=Rebaseline.RB)])
        shutil.move(str(self.v.root / ABC), str(self.v.root / "projects/XPL Example Plus"))
        self.v.run("recode", "ABC", "XPL", "--name", "Example Plus", "--folder", "XPL Example Plus", script="propose.py")
        self.v.run("apply")
        m = self.ms("XPL-M-002")
        self.assertEqual((m["owner"], m["state"], m["state_set"], m["client_visible"], m["baseline_date"]),
                         ("Sara Example", "At risk", 1, 1, iso(80)))
        sys.path.insert(0, str(self.v.root / "system/scripts"))
        import tracker
        with closing(tracker.connect(self.v.root / "system/tracker.db")) as con:
            self.assertEqual(tracker.milestone_original(con, "XPL-M-002"), iso(60))
            ev = tracker.milestone_events(con, ["XPL-M-002"])["XPL-M-002"]
            self.assertIn(("baseline_date", iso(60), iso(80)), [(e["field"], e["old"], e["new"]) for e in ev])
        self.assertEqual([d["id"][:4] for d in self.v.rows("decisions")], ["XPL-"])


class Excel(Case):
    def test_export_shows_the_new_columns_and_import_stages_them(self):
        self.apply("x", [upd("ABC-M-001", owner="Sara Example", state="At risk", state_reason="Review open")])
        wb = load_workbook(self.v.root / "tracker-view.xlsx")
        ws = wb["Milestones"]
        h = [c.value for c in ws[1]]
        for col in ("Owner", "State", "State set by you", "State reason", "Achieved on", "Client-visible", "Client wording", "Internal note"):
            self.assertIn(col, h)
        row = {h[i]: c.value for i, c in enumerate(ws[2])}
        self.assertEqual((row["Owner"], row["State"], row["State set by you"]), ("Sara Example", "At risk", "Yes"))
        ws.cell(row=2, column=h.index("State") + 1).value = "Slipped"
        ws.cell(row=2, column=h.index("Client-visible") + 1).value = "Yes"
        ws.cell(row=2, column=h.index("Client wording") + 1).value = "Expected a week later than agreed."
        ws.cell(row=2, column=h.index("Baseline") + 1).value = date.fromisoformat(iso(45))     # a rebaseline: ignored, with a note
        wb.save(self.v.root / "tracker-view.xlsx")
        out = self.v.run("import", script="excel_sync.py").stdout
        self.assertIn("is a rebaseline", out)
        p = sorted((self.v.root / "system/staging").glob("proposed-*-excel*.json"))[-1]
        data = json.loads(p.read_text())
        self.assertEqual(data["items"][0]["changes"], {"state": "Slipped", "client_visible": 1,
                                                       "client_message": "Expected a week later than agreed."})
        data["items"][0]["decision"] = "approved"
        p.write_text(json.dumps(data))
        self.v.run("apply")
        m = self.ms()
        self.assertEqual((m["state"], m["state_set"], m["client_visible"], m["baseline_date"]), ("Slipped", 1, 1, iso(30)))


class Reports(Case):
    """One vocabulary (the five states) in the weekly report, the client report, the one-pager and the status facts, and the client
    filter (client-visible milestones with their wording only, in the client's words)."""

    def setUp(self):
        super().setUp()
        self.apply("r", [upd("ABC-M-001", n=1, forecast_date=iso(37), client_visible=1, client_message="Testing a week later than agreed."),
                         milestone(2, "Internal security review", iso(20)),
                         milestone(3, "Kick-off", iso(-20), client_visible=1, client_message="Started as planned.")])
        self.apply("d", [upd("ABC-M-004", status="Done", actual_date=iso(-21), state="Achieved")])

    def data(self, *args):
        return json.loads(self.v.run(*args, script="report.py").stdout)

    def test_the_weekly_report_uses_the_five_states(self):
        ms = {m["name"]: m for m in self.data("weekly-data", "ABC")["milestones"]}
        self.assertEqual((ms["UAT complete"]["state"], ms["Go-live"]["state"], ms["Internal security review"]["state"],
                          ms["Kick-off"]["state"]), ("Slipped", "On track", "On track", "Achieved"))

    def test_the_client_report_holds_client_visible_milestones_in_the_clients_words(self):
        ms = self.data("client-data", "ABC", "--since", iso(-30))["milestones"]
        self.assertEqual([(m["name"], m["state"]) for m in ms], [("Kick-off", "Achieved"), ("UAT complete", "Expected later")])
        self.assertIn("expected", next(m for m in ms if m["name"] == "UAT complete")["shown_date"])

    def test_the_one_pager_internal_and_client_safe(self):
        internal = self.data("steerco-data", "ABC")
        client = self.data("steerco-data", "ABC", "--version", "client")
        flat = lambda d: json.dumps(d)  # noqa: E731
        self.assertIn("Internal security review", flat(internal))
        self.assertNotIn("Internal security review", flat(client))
        self.assertNotIn("Delayed", flat(internal) + flat(client))
        self.assertNotIn("Complete\"", flat(internal) + flat(client))
        self.assertIn("Expected later", flat(client))
        self.assertIn("Slipped", flat(internal))

    def test_the_internal_note_never_reaches_client_output(self):
        secret = "Vendor contract dispute, do not share"
        self.apply("n", [upd("ABC-M-001", internal_note=secret), upd("ABC-M-004", n=2, internal_note=secret)])
        client = [self.data("client-data", "ABC", "--since", iso(-30)), self.data("steerco-data", "ABC", "--version", "client")]
        for d in client:
            self.assertNotIn(secret, json.dumps(d))
        r = self.v.run("facts", "ABC", "--audience", "client", "--from", iso(-30), "--to", iso(0), script="status_facts.py")
        self.assertNotIn(secret, r.stdout)
        ctx = self.v.run("current-context", "ABC", "--audience", "client", script="report.py", ok=False).stdout
        self.assertNotIn(secret, ctx)
        spec = self.v.root / "system/staging/drafts/cr.json"
        spec.parent.mkdir(parents=True, exist_ok=True)
        spec.write_text(json.dumps({"code": "ABC", "period_start": iso(-30), "summary": ["A quiet week."], "force": True}))
        self.v.run("client-stage", str(spec), script="report.py")
        draft = next((self.v.root / "system/staging/drafts").glob("*Client report*.html"))   # v2.5.1: the hub's own, no Word file
        self.assertNotIn(secret, draft.read_text(encoding="utf-8"))               # neither its page nor its data block
        self.assertNotIn(secret, __import__("test_reports").client_draft_text(self, self.v))
        self.assertEqual(self.ms()["internal_note"], secret)                     # kept, for Movement

    def test_status_facts_for_a_client_leave_internal_milestones_out(self):
        r = self.v.run("facts", "ABC", "--audience", "client", "--from", iso(-30), "--to", iso(0), script="status_facts.py")
        self.assertNotIn("Internal security review", r.stdout)


class Migration(unittest.TestCase):
    """The Missed migration on a copy of a populated database: every row kept, Missed becomes Planned with state Overdue,
    logged as one apply; a second run does nothing; the migration can't be undone."""

    def setUp(self):
        self.v = Vault()
        populate(self.v)
        self.db = self.v.root / "system/tracker.db"
        with closing(sqlite3.connect(self.db)) as con:
            con.execute("PRAGMA foreign_keys = OFF")
            rows = con.execute(f"SELECT {OLD_COLS} FROM milestones").fetchall()
            self.assertGreaterEqual(len(rows), 2)
            con.execute("ALTER TABLE milestones RENAME TO milestones_new_tmp")
            con.execute(OLD_MILESTONES)
            con.executemany(f"INSERT INTO milestones ({OLD_COLS}) VALUES ({', '.join('?' * 11)})", rows)
            con.execute("DROP TABLE milestones_new_tmp")
            self.missed = [r[0] for r in rows[:2]]
            con.executemany("UPDATE milestones SET status = 'Missed' WHERE id = ?", [(i,) for i in self.missed])
            con.commit()
            self.counts = self.counts_now(con)
            self.before = con.execute(f"SELECT {OLD_COLS} FROM milestones ORDER BY id").fetchall()

    def tearDown(self):
        self.v.cleanup()

    def counts_now(self, con):
        tables = [r[0] for r in con.execute("SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'vault_search%' "
                                            "AND name NOT LIKE 'sqlite_%'")]
        return {t: con.execute(f"SELECT count(*) FROM {t}").fetchone()[0] for t in tables if t not in ("change_log", "applies", "item_events")}   # item_events: the work-item-history table, new with that schema, not touched by this migration

    def test_missed_becomes_overdue_rows_are_kept_and_a_second_run_does_nothing(self):
        out = self.v.run("migrate").stdout
        self.assertIn("Missed is retired", out)
        for i in self.missed:
            self.assertIn(i, out)
        with closing(sqlite3.connect(self.db)) as con:
            con.row_factory = sqlite3.Row
            self.assertEqual(self.counts_now(con), self.counts)                              # every row in every table kept
            after = {r["id"]: dict(r) for r in con.execute("SELECT * FROM milestones")}
            for r in self.before:
                a = after[r[0]]
                self.assertEqual((a["project"], a["name"], a["baseline_date"], a["forecast_date"], a["percent_complete"]),
                                 (r[1], r[2], r[4], r[5], r[8]))
                if r[0] in self.missed:
                    self.assertEqual((a["status"], a["state"], a["state_set"]), ("Planned", "Overdue", 1))
                else:
                    self.assertEqual((a["status"], a["state"], a["state_set"]), (r[6], None, 0))
            log = con.execute("SELECT * FROM change_log WHERE via = 'migrate' ORDER BY rowid").fetchall()
            self.assertEqual({(x["item_id"], x["field"], x["old_value"], x["new_value"]) for x in log if x["field"] in ("status", "state")},
                             {(i, f, o, n) for i in self.missed for f, o, n in (("status", "Missed", "Planned"), ("state", None, "Overdue"))})
            self.assertEqual(len({x["apply_id"] for x in log}), 1)
            n_log = con.execute("SELECT count(*) FROM change_log").fetchone()[0]
        with closing(sqlite3.connect(self.db)) as con:
            self.assertIn("internal_note", [r[1] for r in con.execute("PRAGMA table_info(milestones)")])
        out = self.v.run("migrate").stdout                                                  # a second run: nothing changes
        self.assertIn("Schema is up to date", out)
        with closing(sqlite3.connect(self.db)) as con:
            self.assertEqual(con.execute("SELECT count(*) FROM change_log").fetchone()[0], n_log)
        r = self.v.run("undo", ok=False)                                                    # Missed can't come back
        self.assertIn("can't be undone", r.stdout)
        with closing(sqlite3.connect(self.db)) as con:
            self.assertEqual(con.execute("SELECT count(*) FROM milestones WHERE state = 'Overdue'").fetchone()[0], len(self.missed))

    def test_activity_lists_the_migration(self):
        self.v.run("migrate")
        r = subprocess.run([sys.executable, str(self.v.root / "system/scripts/build_dashboard.py")], cwd=self.v.root,
                           capture_output=True, text=True)
        self.assertEqual(r.returncode, 0, r.stderr)
        html = (self.v.root / "dashboard.html").read_text()
        self.assertIn("Missed is retired", html)


if __name__ == "__main__":
    unittest.main()
