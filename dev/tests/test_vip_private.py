"""BL-288 part 1: VIP is private to Alan. With a VIP set in the people register, nothing creaite writes outside the
dashboard says "VIP": the weekly report, the client report, the steerco one-pager (both versions), a one-off Word
document, a prep page and the Excel export. ask creaite's prompts carry the rule, and the skills don't ask for the label.
VIPs keep working (alerts, sorting, the early warning): test_sweep_alerts covers that behaviour.

Run: python3 dev/tests/test_vip_private.py     (quick tier: no Chrome)
"""
import json
import re
import sys
import unittest
import zipfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[1]
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(REPO / "system" / "scripts"))
sys.path.insert(0, str(REPO / "system" / "scripts" / "vendor"))
import test_reports as R  # noqa: E402  (its fixture; its tests don't run here)
import test_steerco as S  # noqa: E402  (its Base fixture; its tests don't run here)
from test_decisions_prep import prep_world  # noqa: E402
from test_tracker import Vault, action  # noqa: E402

WORD = re.compile(r"\bVIP\b")                 # the label, as a word, case-sensitive


def make_vip(v, *names):
    """Mark these people VIPs in the vault's people register (BL-99's VIP column)."""
    rows = "".join(f"| {n} | {n.split()[0].lower()}@example.com | {n.split()[0]} | Example | Lead | | management |\n" for n in names)
    (v.root / "projects/_portfolio/people-register.md").write_text(
        "# People register\n\n## People\n\n| Full name | Email | Aliases | Organisation | Role | Hidden | VIP |\n"
        "|---|---|---|---|---|---|---|\n" + rows, encoding="utf-8")


def docx_text(path):
    with zipfile.ZipFile(path) as z:
        return re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", z.read("word/document.xml").decode("utf-8")))


class NoLabel:
    def assertNoVip(self, text, what):
        m = WORD.search(text)
        self.assertIsNone(m, f"{what} says VIP: …{text[max(0, m.start() - 80):m.end() + 80] if m else ''}…")


class WeeklyAndClient(NoLabel, unittest.TestCase):
    def setUp(self):
        self.v = R.ReportVault()
        R.populate(self.v)
        make_vip(self.v, "Alan Carter")                     # owns the overdue ABC-A-001: the vip_overdue early warning fires

    def tearDown(self):
        self.v.cleanup()

    def test_weekly_data_and_report_never_say_vip(self):
        data = self.v.at("weekly-data", "ABC", script="report.py").stdout
        d = json.loads(data[data.index("{"):])
        self.assertEqual(d["vip"]["owners"], [["Alan Carter", "management"]])                  # the signal is still there
        self.assertTrue(any("owned by Alan Carter" in w for w in d["early_warnings"]), d["early_warnings"])
        self.assertNoVip(json.dumps(d["early_warnings"]), "the weekly early warnings")
        f = self.v.root / "system/staging/drafts/weekly.json"
        f.write_text(json.dumps({"code": "ABC", "health": "At risk", "health_reason": "UAT start slipped 14 days",
                                 "summary": ["UAT start has slipped two weeks."], "key_changes": "UAT start moved",
                                 "current_state": "Recovery plan agreed.", "progress": ["Plan revised"],
                                 "next_week": ["Escalate data access"], "help_needed": ["ABC-A-001 is overdue"]}))
        out = self.v.at("weekly-stage", str(f), script="report.py").stdout
        self.assertNoVip(out, "weekly-stage's output")
        prop = next((self.v.root / "system/staging").glob("proposed-*-weekly-ABC.json"))
        self.assertNoVip(prop.read_text(), "the staged weekly proposal")
        self.assertNoVip(docx_text(self.v.root / json.loads(prop.read_text())["report"]["draft"]), "the weekly report")

    def test_client_report_never_says_vip(self):
        f = self.v.root / "system/staging/drafts/client.json"
        f.write_text(json.dumps({"code": "ABC", "summary": ["Phase two is being prepared."], "progress": ["Plan revised"],
                                 "next_steps": ["Sign the statement of work"], "input_needed": ["Signature by 16-10"],
                                 "period_start": "2026-10-01"}))
        self.assertNoVip(self.v.at("client-data", "ABC", "--since", "2026-10-01", script="report.py").stdout, "client-data")
        self.assertNoVip(self.v.at("client-stage", str(f), script="report.py").stdout, "client-stage's output")
        prop = next((self.v.root / "system/staging").glob("proposed-*-client-ABC.json"))
        self.assertNoVip(docx_text(self.v.root / json.loads(prop.read_text())["report"]["draft"]), "the client report")

    def test_a_one_off_document_never_says_vip(self):
        f = self.v.root / "system/staging/drafts/adhoc.json"
        f.write_text(json.dumps({"title": "Open actions", "sections": [
            {"heading": "Actions", "tracker": {"table": "actions", "group_by": "owner"}},
            {"heading": "Overdue", "tracker": {"table": "actions", "overdue": True}}]}))
        self.assertNoVip(self.v.at("adhoc", str(f), script="report.py").stdout, "report.py adhoc")
        docs = list((self.v.root / "system/staging/drafts").glob("*Open actions*.docx"))
        self.assertEqual(len(docs), 1)
        text = docx_text(docs[0])
        self.assertIn("Alan Carter", text)
        self.assertNoVip(text, "the one-off document")


class Steerco(NoLabel, S.Base):
    def one_pager(self, version):
        make_vip(self.v, "Alan Carter", "Sara Example")      # Alan owns the overdue ABC-A-005; Sara a decision needed
        self.assertNoVip(json.dumps(self.data("--version", version)), f"steerco-data ({version})")
        self.assertNoVip(self.stage(version=version).stdout, f"steerco-stage ({version})")
        self.assertNoVip(self.draft_text(), f"the {version} one-pager")

    def test_the_internal_one_pager_never_says_vip(self):
        self.one_pager("internal")

    def test_the_client_one_pager_never_says_vip(self):
        self.one_pager("client")


class PrepPage(NoLabel, unittest.TestCase):
    def setUp(self):
        self.v = Vault()
        prep_world(self.v)                                  # Sara Example is invited and owns the late ABC-D decision
        make_vip(self.v, "Sara Example")

    def tearDown(self):
        self.v.cleanup()

    def test_the_prep_page_and_its_research_prompt_never_say_vip(self):
        import prep_page
        r = self.v.run("context", "09:00|ABC weekly", script="prep_page.py")
        ctx = json.loads(r.stdout[r.stdout.index("{"):])
        whys = [i["why"] for i in ctx["items"] if i.get("why")]
        self.assertIn("Sara Example is invited", " ".join(whys))                              # the VIP still ranks, unlabelled
        self.assertNoVip(json.dumps(ctx), "the prep context")
        self.assertNoVip(prep_page.prompt(ctx), "the prep research prompt")
        self.assertNoVip(prep_page.render(ctx, {}), "the prep page")


class Excel(NoLabel, unittest.TestCase):
    def setUp(self):
        self.v = Vault()
        self.v.run("init")
        make_vip(self.v, "Omar Ali")
        self.v.propose("seed", [action(1, owner="Omar Ali", due_date="2026-10-01")])
        self.v.run("apply")
        self.v.propose("pending", [{**action(2, owner="Omar Ali", description="Reply to Omar"), "flags": ["vip"],
                                    "decision": "pending"}])

    def tearDown(self):
        self.v.cleanup()

    def test_the_export_has_no_vip_column_label_or_flag(self):
        from openpyxl import load_workbook
        self.v.run("export", script="excel_sync.py")
        wb = load_workbook(self.v.root / "tracker-view.xlsx")
        cells = [str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value is not None]
        self.assertTrue(any("Reply to Omar" in c for c in cells))                              # the VIP's proposal is listed
        hits = [c for c in cells if re.search(r"\bvip\b", c, re.I)]
        self.assertEqual(hits, [], "the Excel export names a VIP")


class ThreadCards(NoLabel, unittest.TestCase):
    """BL-342 Part 1: a thread's card (type, ask, due, owners, keys, state), Handled and Log on the matched item never say VIP,
    even when Claude's reply does; the route handlers' own text and the page data for the card are clean."""

    def test_the_card_the_page_data_the_routes_and_the_logged_update_never_say_vip(self):
        import os
        import subprocess
        from test_sweep_alerts import run, vip_file, vip_vault
        v = vip_vault()
        try:
            found = [{"kind": "teams", "conversation": "19:c@thread.v2", "title": "Go-live (VIP)", "from": "Mia Hale <mia@client.example>",
                      "people": ["Mia Hale <mia@client.example>", "Omar Ali <omar@company.example>"], "last": "2026-10-08 08:40",
                      "link": "https://teams.microsoft.com/l/chat/19:c", "summary": "A VIP asks for the date [VIP]",
                      "type": "action", "ask": "Tell the VIP the date", "due": "2026-10-20", "owners": ["VIP Mia Hale", "you"],
                      "items": [{"type": "action", "text": "Send the date to the VIP", "owner": "Omar Ali"}]}]
            f = v.root / "system/staging/drafts/found.json"
            f.write_text(json.dumps(found))
            run(v, "sweep.py", "vip", "record", str(f), "--found", "check", now="2026-10-08 09:30")
            code = ("import json, sys; sys.path.insert(0, 'system/scripts'); import sweep, build_dashboard\n"
                    "a = sweep.vip_state()['alerts'][0]\nsweep.vip_replied(a['id'], '2026-10-08 09:00')\n"
                    "print(json.dumps(build_dashboard.sweep_alerts({})['vip']['alerts']))\n")
            r = subprocess.run([sys.executable, "-c", code], cwd=v.root, capture_output=True, text=True,
                               env={**os.environ, "CREAITE_NOW": "2026-10-08 09:30"})
            self.assertEqual(r.returncode, 0, r.stderr)
            self.assertNoVip(r.stdout, "the card in the page data")
            a = vip_file(v)["alerts"][0]
            self.assertEqual((a["state"], a["type"], a["owners"]), ("replied", "action", ["Mia Hale", "you"]))
            self.assertNoVip((v.root / "system/staging/vip.json").read_text(), "vip.json")
            self.assertNotIn("keys", json.loads(r.stdout)[0])                                   # the keys stay in vip.json
        finally:
            v.cleanup()

    def test_the_two_route_handlers_say_nothing_of_the_label(self):
        import inspect
        import live_server
        for fn in (live_server.alert_handled_route, live_server.alert_log_route):
            with self.subTest(fn.__name__):
                self.assertNoVip(inspect.getsource(fn), fn.__name__)
        self.assertIn("/api/alerts/handled", live_server.SWEEP_ALERT_ROUTES)
        self.assertIn("/api/alerts/log", live_server.SWEEP_ALERT_ROUTES)


class Prompts(unittest.TestCase):
    def test_ask_creaites_prompts_carry_the_rule(self):
        import live_server
        for name in ("ASK_PROMPT", "CONFIRMED_PROMPT"):
            with self.subTest(name):
                text = getattr(live_server, name)
                self.assertIn("never call or label a person a VIP", text)
                self.assertIn("unlabelled", text)
        self.assertIn("Never use the word VIP in a gist", live_server.VIP_PROMPT)

    def test_the_skills_never_ask_for_the_label(self):
        review = (REPO / ".claude/skills/cx-review/SKILL.md").read_text(encoding="utf-8")
        self.assertNotIn("labelled **VIP**", review)
        weekly = (REPO / ".claude/skills/cx-weekly-report/SKILL.md").read_text(encoding="utf-8")
        self.assertNotIn("VIP messages about ABC", weekly)
        self.assertIn("BL-288", weekly)
        self.assertIn("BL-288", (REPO / "system/creaite-rules.md").read_text(encoding="utf-8"))


if __name__ == "__main__":
    unittest.main()
